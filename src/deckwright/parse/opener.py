"""Разбор шаблона в `TemplateSpec`. Вход слоя parse, один на весь пайплайн.

Модуль только собирает результат из частей и кэширует его. Вся добыча знаний
живёт в соседях: `tokens` — палитра, гарнитуры, шкала и поля;
`patterns` — композиции со слайдов-примеров; `recurring` — логотип и
колонтитул; `fonts` — встроенные шрифты; `semantics` — класс композиции.

Кэш по хэшу файла нужен бюджету времени: разбор колоды на полсотни слайдов с
двумя сотнями картинок занимает секунды, а три варианта вёрстки разбирают один
и тот же шаблон трижды.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.presentation import Presentation as PresentationObject

from deckwright.parse import surfaces
from deckwright.parse import tokens as tokens_mod
from deckwright.parse.bookends import (
    bookend_pattern,
    declares_titles,
    find_bookends,
    structural_bookends,
)
from deckwright.parse.fonts import extract_embedded_fonts
from deckwright.parse.geometry import iter_shapes
from deckwright.parse.patterns import (
    _ALGN,
    is_figure_text,
    mine_slide,
    text_align,
    text_valign,
)
from deckwright.parse.pictures import background_grid, busy_regions, photo_boxes
from deckwright.parse.recurring import find_recurring
from deckwright.parse.semantics import classified
from deckwright.schemas import (
    Box,
    Color,
    Ground,
    LayoutSpec,
    Pattern,
    PatternClass,
    Provenance,
    Slot,
    SlotRole,
    SourceKind,
    TableGrid,
    TemplateSpec,
    TextStyle,
    readable_text_color,
)

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

# Тип плейсхолдера OOXML → роль слота.
_PH_ROLE: dict[str, SlotRole] = {
    "title": SlotRole.TITLE,
    "ctrTitle": SlotRole.TITLE,
    "subTitle": SlotRole.SUBTITLE,
    "body": SlotRole.BODY,
    "obj": SlotRole.BODY,
    "tbl": SlotRole.TABLE,
    "chart": SlotRole.CHART,
    "pic": SlotRole.IMAGE,
    "ftr": SlotRole.FOOTER,
    "sldNum": SlotRole.SLIDE_NUMBER,
    "dt": SlotRole.FOOTER,
}

# Кегль по умолчанию, когда шаблон не сказал ничего: нужен только чтобы
# собрать валидный стиль, реальные значения приходят из шкалы.
FALLBACK_SIZE_PT = 18.0


def file_sha256(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _theme_root(prs: PresentationObject) -> etree._Element | None:
    if not len(prs.slide_masters):
        return None
    for rel in prs.slide_masters[0].part.rels.values():
        if rel.reltype.endswith("/theme"):
            return etree.fromstring(rel.target_part.blob)
    return None


def _theme_font_roles(theme: etree._Element | None) -> dict[str, str]:
    """{major|minor: гарнитура} — для разрешения ссылок `+mj-lt` / `+mn-lt`."""
    if theme is None:
        return {}
    roles: dict[str, str] = {}
    for role, key in (("majorFont", "major"), ("minorFont", "minor")):
        latin = theme.find(f".//{{{A_NS}}}{role}/{{{A_NS}}}latin")
        typeface = latin.get("typeface") if latin is not None else None
        if typeface:
            roles[key] = typeface
    return roles


def _theme_families(theme: etree._Element | None) -> list[str]:
    if theme is None:
        return []
    families: list[str] = []
    for role in ("majorFont", "minorFont"):
        latin = theme.find(f".//{{{A_NS}}}{role}/{{{A_NS}}}latin")
        typeface = latin.get("typeface") if latin is not None else None
        if typeface and typeface not in families:
            families.append(typeface)
    return families


def _shape_tree(element: etree._Element) -> etree._Element | None:
    return element.find(f".//{{{P_NS}}}spTree")


def _placeholder_role(element: etree._Element) -> SlotRole:
    ph = element.findall(f".//{{{P_NS}}}ph")
    if not ph:
        return SlotRole.UNKNOWN
    return _PH_ROLE.get(ph[0].get("type", "body"), SlotRole.BODY)


def _slot_size_pt(element: etree._Element) -> float | None:
    """Кегль, которым шаблон набирает этот плейсхолдер.

    Позиция в шкале — плохая замена: верх шкалы у шаблона занят обложечными
    размерами (у `vk_education` это 60 pt), и считать по ним бюджет рабочего
    заголовка значит получить двадцать символов вместо полусотни.
    """
    sizes = [
        int(node.get("sz")) / 100
        for node in element.iter()
        if etree.QName(node).localname in ("defRPr", "rPr") and node.get("sz")
    ]
    return max(sizes) if sizes else None


def _text_color(
    element: etree._Element, theme: dict[str, str], clr_map: dict[str, str]
) -> Color | None:
    """Цвет, которым сам шаблон пишет текст в этой фигуре.

    Надёжнее вывода по яркости фона: шаблон уже принял решение с учётом
    градиентов и фоновых картинок, о которых парсер ничего не знает.

    Цвет — большинства символов абзаца основного текста, а не первого
    попавшегося run'а: доноры выделяют акцентом слово, число или
    подзаголовок карточки, и цвет выделения становился цветом всего места —
    синий основной текст на `vk_tech` и `vk_workspace`. Символы без своего
    цвета пишутся цветом абзаца, затем списка стилей фигуры; нет и его —
    цвет наследуется (None), и решает звено выше.
    """
    shared = _declared_color(element.find(f".//{{{A_NS}}}lstStyle"), theme, clr_map)
    # По абзацам: цвет большинства символов абзаца и его кегль.
    paragraphs: list[tuple[float, Color | None]] = []
    for paragraph in element.iter(f"{{{A_NS}}}p"):
        own = _declared_color(paragraph.find(f"{{{A_NS}}}pPr"), theme, clr_map) or shared
        counts: dict[str | None, tuple[int, Color | None]] = {}
        size = 0.0
        for run in paragraph.findall(f"{{{A_NS}}}r"):
            text = (run.findtext(f"{{{A_NS}}}t") or "").strip()
            if not text:
                continue
            props = run.find(f"{{{A_NS}}}rPr")
            color = _declared_color(props, theme, clr_map) or own
            key = color.rgb if color is not None else None
            number, _ = counts.get(key, (0, color))
            counts[key] = (number + len(text), color)
            if props is not None and props.get("sz"):
                size = max(size, int(props.get("sz")) / 100)
        if counts:
            paragraphs.append((size, max(counts.values(), key=lambda item: item[0])[1]))
    if not paragraphs:
        # Пустая рамка (плейсхолдер макета): цвет из её объявлений.
        for node in element.iter():
            if etree.QName(node).localname in ("defRPr", "rPr", "endParaRPr"):
                color = tokens_mod.resolve_color(node, theme, clr_map)
                if color is not None:
                    return color
        return None
    # Абзацы разного цвета — подзаголовок карточки и её текст («Заголовок»
    # синим, «Текст» белым на `vk_workspace`). Цвет текста — у абзаца
    # основного текста: самого мелкого, при равных — последнего.
    known = [size for size, _ in paragraphs if size > 0]
    smallest = min(known) if known else 0.0
    candidates = [color for size, color in paragraphs if size in (smallest, 0.0)]
    return candidates[-1] if candidates else paragraphs[-1][1]


def _declared_color(
    node: etree._Element | None, theme: dict[str, str], clr_map: dict[str, str]
) -> Color | None:
    """Цвет заливки текста, объявленный прямо в этом узле (`rPr`, `pPr`, `lstStyle`).

    Только своя `a:solidFill` текста: цвета подчёркивания, тени и обводки
    к цвету букв отношения не имеют.
    """
    if node is None:
        return None
    if etree.QName(node).localname in ("pPr",):
        node = node.find(f"{{{A_NS}}}defRPr")
    elif etree.QName(node).localname == "lstStyle":
        level = node.find(f"{{{A_NS}}}lvl1pPr")
        node = level.find(f"{{{A_NS}}}defRPr") if level is not None else None
    if node is None:
        return None
    fill = node.find(f"{{{A_NS}}}solidFill")
    return tokens_mod.resolve_color(fill, theme, clr_map) if fill is not None else None


# Какой раздел `p:txStyles` мастера отвечает за роль плейсхолдера. Так же
# разрешает наследование и сам PowerPoint: кегль, не объявленный в layout'е,
# берётся из стиля мастера, а не выдумывается.
_MASTER_STYLE_BY_ROLE: dict[SlotRole, str] = {
    SlotRole.TITLE: "titleStyle",
    SlotRole.SUBTITLE: "bodyStyle",
    SlotRole.BODY: "bodyStyle",
}


def _master_sizes_pt(master) -> dict[str, float]:
    """Кегли первого уровня из `p:txStyles` мастера.

    Зачем. Большинство layout'ов кегль заголовка не объявляют вовсе — он
    наследуется. Раньше на их месте вставлялась середина типографической
    шкалы, и у чужого шаблона заголовок оказывался набран двадцатым кеглем
    вместо сорок четвёртого: бюджет заголовка вырастал вчетверо и просил у
    модели абзац там, где рамка держит строку.
    """
    tx_styles = master._element.find(f"{{{P_NS}}}txStyles")
    if tx_styles is None:
        return {}
    sizes: dict[str, float] = {}
    for style_name in ("titleStyle", "bodyStyle", "otherStyle"):
        node = tx_styles.find(f"{{{P_NS}}}{style_name}")
        if node is None:
            continue
        level = node.find(f"{{{A_NS}}}lvl1pPr")
        def_rpr = level.find(f"{{{A_NS}}}defRPr") if level is not None else None
        if def_rpr is not None and def_rpr.get("sz"):
            sizes[style_name] = int(def_rpr.get("sz")) / 100
    return sizes


def _master_color(
    master, element: etree._Element, theme: dict[str, str], clr_map: dict[str, str]
) -> Color | None:
    """Цвет текста, который плейсхолдер наследует от мастера.

    Та же цепочка, что у кегля: плейсхолдер мастера того же типа, затем
    раздел `p:txStyles`. Раньше при молчании layout'а шёл цвет, «читаемый
    по фону» из палитры, — не тот, которым шаблон пишет.
    """
    ph = element.find(f".//{{{P_NS}}}ph")
    wanted = ph.get("type", "body") if ph is not None else "body"
    for shape in master.placeholders:
        node = shape._element.find(f".//{{{P_NS}}}ph")
        kind = node.get("type", "body") if node is not None else "body"
        if _PH_ROLE.get(kind) is _PH_ROLE.get(wanted):
            color = _text_color(shape._element, theme, clr_map)
            if color is not None:
                return color
            break
    style_name = _MASTER_STYLE_BY_ROLE.get(_PH_ROLE.get(wanted), "otherStyle")
    return _master_style_color(master, style_name, theme, clr_map)


def _master_align(master, element: etree._Element):
    """Выравнивание, которое плейсхолдер наследует от мастера.

    Цепочка та же, что у цвета и кегля: плейсхолдер мастера того же типа,
    затем первый уровень раздела `p:txStyles`.
    """
    ph = element.find(f".//{{{P_NS}}}ph")
    wanted = ph.get("type", "body") if ph is not None else "body"
    for shape in master.placeholders:
        node = shape._element.find(f".//{{{P_NS}}}ph")
        kind = node.get("type", "body") if node is not None else "body"
        if _PH_ROLE.get(kind) is _PH_ROLE.get(wanted):
            align = text_align(shape._element)
            if align is not None:
                return align
            break
    tx_styles = master._element.find(f"{{{P_NS}}}txStyles")
    style_name = _MASTER_STYLE_BY_ROLE.get(_PH_ROLE.get(wanted), "otherStyle")
    node = tx_styles.find(f"{{{P_NS}}}{style_name}") if tx_styles is not None else None
    level = node.find(f"{{{A_NS}}}lvl1pPr") if node is not None else None
    return _ALGN.get(level.get("algn", "")) if level is not None else None


def _master_style_color(
    master, style_name: str, theme: dict[str, str], clr_map: dict[str, str]
) -> Color | None:
    """Цвет первого уровня раздела `p:txStyles` мастера."""
    tx_styles = master._element.find(f"{{{P_NS}}}txStyles")
    node = tx_styles.find(f"{{{P_NS}}}{style_name}") if tx_styles is not None else None
    level = node.find(f"{{{A_NS}}}lvl1pPr") if node is not None else None
    def_rpr = level.find(f"{{{A_NS}}}defRPr") if level is not None else None
    return tokens_mod.resolve_color(def_rpr, theme, clr_map) if def_rpr is not None else None


def _typeface_in(element: etree._Element) -> str | None:
    """Гарнитура, объявленная внутри фигуры. None — не объявлена нигде.

    Сначала смотрим прогоны текста: ими набрано то, что видит читатель.
    Списочные стили и `defRPr` идут следом — они задают умолчание фигуры.
    """
    runs = [
        node
        for run in element.iter(f"{{{A_NS}}}r")
        for node in run.iter(f"{{{A_NS}}}latin")
    ]
    for node in runs or list(element.iter(f"{{{A_NS}}}latin")):
        face = node.get("typeface")
        if face:
            return face
    return None


def _resolved_typeface(
    element: etree._Element, theme_fonts: dict[str, str]
) -> str | None:
    """Гарнитура фигуры с разрешённой ссылкой на тему (`+mj-lt`, `+mn-lt`)."""
    return tokens_mod.resolve_typeface(_typeface_in(element), theme_fonts)


def _master_style_typeface(
    master, style_name: str, theme_fonts: dict[str, str]
) -> str | None:
    """Гарнитура первого уровня раздела `p:txStyles` мастера."""
    tx_styles = master._element.find(f"{{{P_NS}}}txStyles")
    node = tx_styles.find(f"{{{P_NS}}}{style_name}") if tx_styles is not None else None
    level = node.find(f"{{{A_NS}}}lvl1pPr") if node is not None else None
    def_rpr = level.find(f"{{{A_NS}}}defRPr") if level is not None else None
    face = def_rpr.find(f"{{{A_NS}}}latin") if def_rpr is not None else None
    return tokens_mod.resolve_typeface(
        face.get("typeface") if face is not None else None, theme_fonts
    )


def _master_placeholder_typeface(
    master, element: etree._Element, theme_fonts: dict[str, str]
) -> str | None:
    """Гарнитура плейсхолдера мастера того же типа — звено наследования OOXML."""
    ph = element.find(f".//{{{P_NS}}}ph")
    wanted = ph.get("type", "body") if ph is not None else "body"
    for shape in master.placeholders:
        node = shape._element.find(f".//{{{P_NS}}}ph")
        kind = node.get("type", "body") if node is not None else "body"
        if _PH_ROLE.get(kind) is not _PH_ROLE.get(wanted):
            continue
        face = _resolved_typeface(shape._element, theme_fonts)
        if face:
            return face
    return None


def _placeholder_typeface(
    layout, element: etree._Element, role: SlotRole, theme_fonts: dict[str, str]
) -> str | None:
    """Чем на самом деле будет набран этот плейсхолдер.

    Цепочка наследования OOXML, звено за звеном: сам плейсхолдер → тот же
    плейсхолдер мастера → раздел `p:txStyles` мастера → шрифт темы по роли.
    Последнее звено — не догадка, а правило формата: заголовок берёт
    `majorFont`, остальной текст `minorFont`. Без него заголовки колоды
    мерились гарнитурой основного текста, а рисовались заголовочной
    (`Calibri` против `Calibri Light` на `finansy`).
    """
    master = layout.slide_master
    found = (
        _resolved_typeface(element, theme_fonts)
        or _master_placeholder_typeface(master, element, theme_fonts)
        or _master_style_typeface(
            master, _MASTER_STYLE_BY_ROLE.get(role, "otherStyle"), theme_fonts
        )
    )
    if found:
        return found
    return theme_fonts.get("major" if role is SlotRole.TITLE else "minor")


def _weight_in(element: etree._Element) -> tuple[bool | None, bool | None]:
    """Начертание прогонов фигуры: (полужирное, курсив). None — не объявлено.

    Не объявлено — значит наследуется, и подменять это `False` нельзя:
    заголовок, которому полужирность задаёт мастер, стал бы обычным.
    """
    marked = [
        run.find(f"{{{A_NS}}}rPr")
        for run in element.iter(f"{{{A_NS}}}r")
    ]
    marked = [node for node in marked if node is not None]
    if not marked:
        return None, None
    bold = [node.get("b") for node in marked if node.get("b") is not None]
    italic = [node.get("i") for node in marked if node.get("i") is not None]
    return (
        (sum(1 for value in bold if value == "1") * 2 >= len(marked)) if bold else None,
        (sum(1 for value in italic if value == "1") * 2 >= len(marked))
        if italic
        else None,
    )


def _master_placeholder_weight(master, element: etree._Element, role) -> bool | None:
    """Полужирность, которую плейсхолдер наследует от мастера.

    Заголовки большинства шаблонов объявлены полужирными не на слайде, а в
    `p:txStyles` мастера. Без этого звена фиттер мерил заголовок обычным
    начертанием, а рисовался он полужирным — на 7-8 % шире.
    """
    ph = element.find(f".//{{{P_NS}}}ph")
    wanted = ph.get("type", "body") if ph is not None else "body"
    for shape in master.placeholders:
        node = shape._element.find(f".//{{{P_NS}}}ph")
        kind = node.get("type", "body") if node is not None else "body"
        if _PH_ROLE.get(kind) is not _PH_ROLE.get(wanted):
            continue
        bold, _ = _weight_in(shape._element)
        if bold is not None:
            return bold
        break
    tx_styles = master._element.find(f"{{{P_NS}}}txStyles")
    name = _MASTER_STYLE_BY_ROLE.get(role, "otherStyle")
    node = tx_styles.find(f"{{{P_NS}}}{name}") if tx_styles is not None else None
    level = node.find(f"{{{A_NS}}}lvl1pPr") if node is not None else None
    def_rpr = level.find(f"{{{A_NS}}}defRPr") if level is not None else None
    value = def_rpr.get("b") if def_rpr is not None else None
    return value == "1" if value is not None else None


def _master_placeholder_size_pt(master, element: etree._Element) -> float | None:
    """Кегль плейсхолдера мастера того же типа — звено наследования OOXML.

    Цепочка такая: плейсхолдер слайда → макета → **плейсхолдера мастера** →
    `p:txStyles` мастера. Среднее звено раньше пропускалось, и у
    `vk_education` (экспорт из Google Slides) заголовок получал 14 pt из
    `titleStyle`, хотя плейсхолдер заголовка в мастере набран 36 pt — так
    он и рисуется. Бюджет длины заголовка при этом завышался вдвое с лишним.
    Берётся первый уровень: он и есть кегль самого заголовка.
    """
    ph = element.find(f".//{{{P_NS}}}ph")
    wanted = ph.get("type", "body") if ph is not None else "body"
    for shape in master.placeholders:
        node = shape._element.find(f".//{{{P_NS}}}ph")
        kind = node.get("type", "body") if node is not None else "body"
        if _PH_ROLE.get(kind) is not _PH_ROLE.get(wanted):
            continue
        level = shape._element.find(f".//{{{A_NS}}}lvl1pPr/{{{A_NS}}}defRPr")
        if level is not None and level.get("sz"):
            return int(level.get("sz")) / 100
        return _slot_size_pt(shape._element)
    return None


def _layout_slots(
    layout,
    style: TextStyle,
    theme,
    clr_map,
    fallback: Color,
    master_sizes: dict[str, float],
    theme_fonts: dict[str, str] | None = None,
) -> list[Slot]:
    slots: list[Slot] = []
    for shape in layout.placeholders:
        if None in (shape.left, shape.top, shape.width, shape.height):
            continue
        if shape.width <= 0 or shape.height <= 0:
            continue
        fmt = shape.placeholder_format
        color = (
            _text_color(shape._element, theme, clr_map)
            or _master_color(layout.slide_master, shape._element, theme, clr_map)
            or fallback
        )
        role = _placeholder_role(shape._element)
        size = (
            _slot_size_pt(shape._element)
            or _master_placeholder_size_pt(layout.slide_master, shape._element)
            or master_sizes.get(_MASTER_STYLE_BY_ROLE.get(role, "otherStyle"))
        )
        updates = {"color": color}
        if size:
            updates["size_pt"] = size
        # Гарнитура и начертание — свои у каждого места. Раньше всем
        # доставалась одна гарнитура колоды, и заголовок, который шаблон
        # набирает заголовочным шрифтом темы, мерился основным.
        family = _placeholder_typeface(layout, shape._element, role, theme_fonts or {})
        if family:
            updates["font_family"] = family
        bold, italic = _weight_in(shape._element)
        if bold is None:
            bold = _master_placeholder_weight(layout.slide_master, shape._element, role)
        if bold is not None:
            updates["bold"] = bold
        if italic is not None:
            updates["italic"] = italic
        align = text_align(shape._element) or _master_align(layout.slide_master, shape._element)
        if align is not None:
            updates["align"] = align
        valign = text_valign(shape._element)
        if valign is not None:
            updates["valign"] = valign
        slots.append(
            Slot(
                id=f"ph{fmt.idx}",
                role=role,
                box=Box(x=shape.left, y=shape.top, w=shape.width, h=shape.height),
                style=style.model_copy(update=updates),
                placeholder_text=shape.text_frame.text if shape.has_text_frame else "",
                ph_idx=fmt.idx,
                provenance=Provenance(kind=SourceKind.LAYOUT, ref=layout.name),
            )
        )
    return slots


def _parse(path: Path, font_dir: Path | None) -> TemplateSpec:
    prs = Presentation(str(path))
    slide_w, slide_h = prs.slide_width, prs.slide_height
    theme_root = _theme_root(prs)
    theme = tokens_mod.theme_colors(theme_root)
    theme_fonts = _theme_font_roles(theme_root)
    warnings: list[str] = []

    # ── Статистика по слайдам, layout'ам и мастерам ──────────────────────────
    usage = tokens_mod.Usage()
    for master in prs.slide_masters:
        clr_map = tokens_mod.color_map(master._element)
        master_sizes = _master_sizes_pt(master)
        trees = [_shape_tree(master._element)]
        trees += [_shape_tree(layout._element) for layout in master.slide_layouts]
        tokens_mod.collect(
            [t for t in trees if t is not None],
            theme,
            clr_map,
            slide_w,
            slide_h,
            usage,
            on_slide=False,
            theme_fonts=theme_fonts,
        )
    primary_map = (
        tokens_mod.color_map(prs.slide_masters[0]._element) if len(prs.slide_masters) else {}
    )
    tokens_mod.collect(
        [slide.shapes._spTree for slide in prs.slides],
        theme,
        primary_map,
        slide_w,
        slide_h,
        usage,
        theme_fonts=theme_fonts,
    )

    # Цвета текста из `p:txStyles` мастеров — тоже цвета шаблона: ими пишется
    # всё, что не объявило своего. Без них палитра не знала цвета, которым
    # шаблон набирает текст по умолчанию, и вёрстка, взяв его, «придумывала».
    for master in prs.slide_masters:
        clr_map = tokens_mod.color_map(master._element)
        for style_name in ("titleStyle", "bodyStyle", "otherStyle"):
            color = _master_style_color(master, style_name, theme, clr_map)
            if color is not None:
                usage.colors[(color.rgb, tokens_mod.ROLE_TEXT)] += 0
                usage.color_counts[(color.rgb, tokens_mod.ROLE_TEXT)] += 1
    palette = tokens_mod.build_palette(usage, theme)
    fonts = tokens_mod.build_fonts(usage, _theme_families(theme_root))
    type_scale = tokens_mod.build_type_scale(usage)
    grid = tokens_mod.build_grid(usage)

    # ── Встроенные шрифты ────────────────────────────────────────────────────
    if font_dir is not None:
        extracted, font_warnings = extract_embedded_fonts(path, font_dir / path.stem)
        warnings.extend(font_warnings)
        by_family = {font.family for font in extracted}
        fonts = [
            token.model_copy(
                update={
                    "embedded": True,
                    "file_path": str(
                        next(f.path for f in extracted if f.family == token.family)
                    ),
                }
            )
            if token.family in by_family
            else token
            for token in fonts
        ]
        missing = [t.family for t in fonts if not t.embedded and t.usage_count > 0]
        if missing:
            warnings.append(
                "шрифты не встроены в шаблон и будут подставлены системой, "
                f"метрики могут разойтись: {', '.join(missing[:5])}"
            )

    base_family = fonts[0].family if fonts else "Arial"
    base_size = type_scale[len(type_scale) // 2] if type_scale else FALLBACK_SIZE_PT
    # Цвет текста, которым шаблон пишет вне плейсхолдеров: `otherStyle`
    # мастера. Самый частый цвет палитры им не был: у `vk_tech` это серый
    # `C4C4C4` подложек, и серым выходил весь текст мест без своего цвета.
    base_color = (
        _master_style_color(prs.slide_masters[0], "otherStyle", theme, primary_map)
        if len(prs.slide_masters)
        else None
    ) or (Color(rgb=palette[0].color.rgb) if palette else Color(rgb="000000"))
    base_style = TextStyle(font_family=base_family, size_pt=base_size, color=base_color)

    # ── Layout'ы ─────────────────────────────────────────────────────────────
    layouts: list[LayoutSpec] = []
    masters: list[str] = []
    layout_ids: dict[int, str] = {}
    # Ограничения макета под слайдом: их надо знать и композиции, а не только
    # самому макету — вёрстка выбирает место именно по композиции.
    obstacles_by_layout: dict[str, list[Box]] = {}
    backdrops_by_layout: dict[str, list[Box]] = {}
    imagery_by_layout: dict[str, list[Box]] = {}
    for m_index, master in enumerate(prs.slide_masters):
        master_id = f"master{m_index + 1}"
        masters.append(master_id)
        clr_map = tokens_mod.color_map(master._element)
        master_bg = tokens_mod.resolve_color(
            master._element.find(f".//{{{P_NS}}}bg"), theme, clr_map
        ) if master._element.find(f".//{{{P_NS}}}bg") is not None else None

        master_tree = _shape_tree(master._element)
        master_backdrop = (
            tokens_mod.backdrop_color(master_tree, slide_w, slide_h, theme, clr_map)
            if master_tree is not None
            else None
        )

        for l_index, layout in enumerate(master.slide_layouts):
            bg_node = layout._element.find(f".//{{{P_NS}}}bg")
            layout_tree = _shape_tree(layout._element)
            background = (
                (
                    tokens_mod.resolve_color(bg_node, theme, clr_map)
                    if bg_node is not None
                    else None
                )
                or (
                    tokens_mod.backdrop_color(layout_tree, slide_w, slide_h, theme, clr_map)
                    if layout_tree is not None
                    else None
                )
                or master_bg
                or master_backdrop
            )
            dark = background is not None and background.luminance < 0.5
            # Шаблон не сказал, каким цветом писать. Спрашиваем его палитру —
            # `#111111` был бы цветом, которого в шаблоне нет, и аудит потом
            # справедливо помечал бы им каждый абзац колоды.
            fallback = readable_text_color(palette, background, dark) or (
                Color(rgb="FFFFFF") if dark else Color(rgb="111111")
            )
            visible = (
                surfaces.collect(
                    layout._element, layout.part, master._element, master.part,
                    slide_w, slide_h,
                )
                if layout_tree is not None
                else []
            )
            # Не просто «поверх чего нельзя писать», а защищённые зоны:
            # логотип из нескольких фигур склеен в одну, рамка во весь
            # слайд сведена к видимым полосам по краям, колонтитулы и
            # номера страниц учтены.
            layout_obstacles = surfaces.protected(visible, slide_w, slide_h)
            layout_backdrops = surfaces.backdrops(visible)
            layout_id = f"{master_id}/layout{l_index + 1}"
            obstacles_by_layout[layout_id] = layout_obstacles
            backdrops_by_layout[layout_id] = layout_backdrops
            imagery_by_layout[layout_id] = surfaces.imagery(visible)
            layout_ids[id(layout._element)] = layout_id
            layouts.append(
                LayoutSpec(
                    id=layout_id,
                    name=layout.name,
                    master_id=master_id,
                    slots=_layout_slots(
                        layout, base_style, theme, clr_map, fallback, master_sizes,
                        theme_fonts,
                    ),
                    # Оформление макета и образца в координатах слайда.
                    # Раньше эти элементы не доезжали до вёрстки вовсе, и
                    # текст ложился поверх картинок шаблона.
                    decor_boxes=layout_obstacles,
                    backdrop_boxes=layout_backdrops,
                    background=background,
                    is_dark=dark,
                )
            )

    if not any(
        slot.role not in (SlotRole.TITLE, SlotRole.FOOTER, SlotRole.SLIDE_NUMBER)
        for layout in layouts
        for slot in layout.slots
    ):
        warnings.append(
            "ни в одном layout'е нет контентных плейсхолдеров: "
            "вёрстка опирается на паттерны со слайдов-примеров"
        )

    # ── Паттерны со слайдов-примеров ─────────────────────────────────────────
    by_layout_id = {layout.id: layout for layout in layouts}
    layout_use: dict[int, int] = {}
    for slide in prs.slides:
        key = id(slide.slide_layout._element)
        layout_use[key] = layout_use.get(key, 0) + 1
    patterns = []
    backdrops: dict[int, Color | None] = {}
    own_backgrounds: dict[int, tuple[Color | None, str | None]] = {}
    titles_declared = declares_titles(list(prs.slides))
    for index, slide in enumerate(prs.slides, start=1):
        layout_id = layout_ids.get(id(slide.slide_layout._element))
        # Фон слайда может перекрывать фон layout'а собственной подложкой.
        own_backdrop = tokens_mod.backdrop_color(
            slide.shapes._spTree, slide_w, slide_h, theme, primary_map
        )
        inherited = by_layout_id.get(layout_id).background if layout_id in by_layout_id else None
        # Собственный фон слайда (`p:bg`) перекрывает фон layout'а и мастера:
        # бланк «Дорожная карта» красит так каждый слайд, а финал — в тёмный.
        bg_node = tokens_mod.own_background(slide._element)
        own_bg = (
            tokens_mod.resolve_color(bg_node, theme, primary_map)
            if bg_node is not None
            else None
        )
        effective = own_backdrop or own_bg or inherited
        pattern = mine_slide(
            slide.shapes._spTree,
            index,
            slide_w,
            slide_h,
            layout_id,
            base_style,
            is_dark=effective is not None and effective.luminance < 0.5,
            inherited_slots=(
                {slot.role: slot for slot in by_layout_id[layout_id].slots}
                if layout_id in by_layout_id
                else None
            ),
            color_of=lambda element: _text_color(element, theme, primary_map),
            theme_fonts=theme_fonts,
        )
        if pattern is not None:
            # Локальный фон каждого слота: на чём лежит текст этого места.
            tree = slide.shapes._spTree
            pattern = pattern.model_copy(
                update={
                    "slots": [
                        slot.model_copy(
                            update={
                                "backdrop": tokens_mod.local_backdrop(
                                    tree, slot.box, theme, primary_map, effective
                                )
                            }
                        )
                        for slot in pattern.slots
                    ],
                    "grounds": _grounds(
                        tree,
                        theme,
                        primary_map,
                        effective,
                        slide_w,
                        slide_h,
                        backgrounds=_background_sources(slide),
                    ),
                    "figure_pictures": _figure_pictures(tree),
                    "background": own_bg,
                    "photo_slots": photo_boxes(
                        tree, slide.part, slide_w, slide_h, illustrations=True
                    ),
                    "background_signature": tokens_mod.background_signature(
                        bg_node, slide.part
                    ),
                    "decor": _decor(tree, pattern, slide_w, slide_h),
                    "layout_obstacles": obstacles_by_layout.get(layout_id, [])
                    + _slide_obstacles(tree, slide.part, pattern, slide_w, slide_h)
                    # Оформление, запечённое в фоновую картинку: логотипы
                    # партнёров ЛЦТ2026 — часть фона, а не фигуры слайда.
                    + _background_decor(slide, slide_w, slide_h),
                    "layout_backdrops": backdrops_by_layout.get(layout_id, [])
                    + surfaces.backdrops(
                        surfaces.collect_tree(tree, slide.part, slide_w, slide_h)
                    ),
                    "baked_items": bool(pattern.repeaters)
                    and _layout_draws_items(slide.slide_layout, layout_use, slide_w, slide_h),
                    "repeaters": [
                        repeater.model_copy(
                            update={
                                "member_backdrops": [
                                    tokens_mod.local_backdrop(
                                        tree,
                                        _text_of_member(repeater, dx, dy),
                                        theme,
                                        primary_map,
                                        effective,
                                    )
                                    for dx, dy in repeater.member_offsets
                                ],
                                "member_aligns": [
                                    _member_align(tree, _text_of_member(repeater, dx, dy))
                                    for dx, dy in repeater.member_offsets
                                ],
                            }
                        )
                        for repeater in pattern.repeaters
                    ],
                }
            )
            pattern = _sequence_ordinals(pattern, tree)
            if not titles_declared:
                pattern = _table_grid(
                    _grown_frames(
                        tree, pattern, slide_w, slide_h,
                        obstacles=obstacles_by_layout.get(layout_id, []),
                    )
                )
            patterns.append(classified(pattern, slide_w, slide_h))
        backdrops[index] = effective
        own_backgrounds[index] = (
            own_bg,
            tokens_mod.background_signature(bg_node, slide.part),
        )

    # ── Обложка и финал ──────────────────────────────────────────────────────
    slides = list(prs.slides)
    cover_index, closing_index = find_bookends(slides, slide_h)
    bookend_ids: dict[PatternClass, str | None] = {}
    for kind, number in ((PatternClass.TITLE, cover_index), (PatternClass.CLOSING, closing_index)):
        bookend_ids[kind] = None
        if number is None:
            continue
        slide = slides[number - 1]
        layout_id = layout_ids.get(id(slide.slide_layout._element))
        backdrop = backdrops.get(number)
        bookend = bookend_pattern(
            slide,
            number,
            kind,
            layout_id,
            by_layout_id[layout_id].slots if layout_id in by_layout_id else [],
            base_style,
            slide_w,
            slide_h,
            is_dark=backdrop is not None and backdrop.luminance < 0.5,
            obstacles=obstacles_by_layout.get(layout_id, []),
            pictures=imagery_by_layout.get(layout_id, []),
        )
        if bookend is not None:
            # Повторители того же слайда — спикеры, контакты: неиспользованный
            # рендер уберёт целиком, с кружком под фото.
            mined = next((p for p in patterns if p.donor_slide_index == number), None)
            if mined is not None:
                # Одиночный спикер, уже входящий в повторитель слайда, второй
                # раз не заводится.
                covered = [frame for r in mined.repeaters for frame in r.member_frames]
                own = [
                    r
                    for r in bookend.repeaters
                    if not any(_overlap(r.item_box, frame) for frame in covered)
                ]
                bookend = bookend.model_copy(update={"repeaters": own + mined.repeaters})
            color, signature = own_backgrounds.get(number, (None, None))
            bookend = bookend.model_copy(
                update={
                    "background": color,
                    "background_signature": signature,
                    "photo_slots": photo_boxes(
                        slide.shapes._spTree, slide.part, slide_w, slide_h
                    ),
                    # Обложка и финал строятся отдельным путём, но оформление
                    # макета им нужно ровно так же: именно на обложке
                    # `vk_tech` подзаголовок и уезжал на картинку макета.
                    "layout_obstacles": obstacles_by_layout.get(layout_id, [])
                    + _slide_obstacles(
                        slide.shapes._spTree, slide.part, bookend, slide_w, slide_h
                    )
                    + _background_decor(slide, slide_w, slide_h),
                    # Фон донора обложке и финалу нужен так же, как всем: без
                    # него цвет текста выбирался по «шаблон светлый», и на
                    # тёмно-фиолетовой обложке ЛЦТ2026 заголовок выходил
                    # чёрным — 2.6:1.
                    "grounds": _grounds(
                        slide.shapes._spTree,
                        theme,
                        primary_map,
                        color,
                        slide_w,
                        slide_h,
                        backgrounds=_background_sources(slide),
                    ),
                    "layout_backdrops": backdrops_by_layout.get(layout_id, [])
                    + surfaces.backdrops(
                        surfaces.collect_tree(
                            slide.shapes._spTree, slide.part, slide_w, slide_h
                        )
                    ),
                }
            )
            patterns.append(bookend)
            bookend_ids[kind] = bookend.id

    if cover_index is None and closing_index is None and not declares_titles(slides):
        cover, closing = structural_bookends(patterns)
        for kind, found in ((PatternClass.TITLE, cover), (PatternClass.CLOSING, closing)):
            if found is None:
                continue
            patterns = [
                _as_bookend(pattern, kind, slide_h) if pattern is found else pattern
                for pattern in patterns
            ]
            bookend_ids[kind] = found.id

    slide_count = len(prs.slides._sldIdLst)
    if slide_count and len(patterns) / slide_count < 0.5:
        warnings.append(
            f"композиции сняты лишь с {len(patterns)} слайдов из {slide_count}: "
            "шаблон беден примерами, вёрстка будет опираться на поля и сетку"
        )

    patterns = _repeated_labels_as_footers(patterns, slide_h)

    # ── Повторяющиеся элементы ───────────────────────────────────────────────
    per_slide: list[list[etree._Element]] = []
    for slide in prs.slides:
        trees = [slide.shapes._spTree]
        for source in (slide.slide_layout, slide.slide_layout.slide_master):
            tree = _shape_tree(source._element)
            if tree is not None:
                trees.append(tree)
        per_slide.append(trees)
    recurring = find_recurring(per_slide, slide_w, slide_h)

    return TemplateSpec(
        template_sha256=file_sha256(path),
        source_name=path.name,
        slide_width_emu=slide_w,
        slide_height_emu=slide_h,
        bullet_indent_emu=_list_indent(prs),
        palette=palette,
        fonts=fonts,
        type_scale_pt=type_scale,
        grid=grid,
        masters=masters,
        layouts=layouts,
        patterns=patterns,
        cover_pattern_id=bookend_ids[PatternClass.TITLE],
        closing_pattern_id=bookend_ids[PatternClass.CLOSING],
        recurring=recurring,
        warnings=warnings,
    )


def _layout_draws_items(layout, layout_use: dict[int, int], slide_w: int, slide_h: int) -> bool:
    """Нарисованы ли элементы композиции в картинке её layout'а.

    Layout, свой у одного слайда, с картинкой крупнее половины слайда: так
    на `vk_tech` сделана сетка 2×2 — карточки и номера «01–04» в фоне.
    """
    if layout_use.get(id(layout._element), 0) > 1:
        return False
    for shape in layout.shapes:
        if shape.is_placeholder or shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        if (shape.width or 0) * (shape.height or 0) >= 0.5 * slide_w * slide_h:
            return True
    return False


# Роли мест, которые не текст: их рамки — не повод считать фигуру подложкой.
_NOT_TEXT_ROLES = frozenset(
    {SlotRole.IMAGE, SlotRole.ICON, SlotRole.CHART, SlotRole.TABLE, SlotRole.DECOR, SlotRole.LOGO}
)


# Нижняя полоса слайда, где живут колонтитулы.
_FOOTER_BAND = 0.8


def _repeated_labels_as_footers(patterns: list[Pattern], slide_h: int) -> list[Pattern]:
    """Одна и та же подпись внизу нескольких слайдов — колонтитул, не место.

    «Презентация создана в Fibonacci» стоит внизу обложки и финала шаблона
    экзаменов. Другого текстового места на обложке нет, и подзаголовок колоды
    садился в эту строку и переполнял её. Надпись, повторённая дословно на
    двух и более слайдах в нижней полосе, — оформление шаблона.
    """
    seen: dict[str, set[int]] = {}
    for pattern in patterns:
        for slot in pattern.slots:
            text = " ".join(slot.placeholder_text.split())
            if text and slot.box.y >= slide_h * _FOOTER_BAND:
                seen.setdefault(text, set()).add(pattern.donor_slide_index)
    repeated = {text for text, slides in seen.items() if len(slides) >= 2}
    if not repeated:
        return patterns

    def footer(slot: Slot) -> Slot:
        text = " ".join(slot.placeholder_text.split())
        if (
            text in repeated
            and slot.box.y >= slide_h * _FOOTER_BAND
            and slot.role not in (SlotRole.SLIDE_NUMBER, SlotRole.TITLE)
        ):
            return slot.model_copy(update={"role": SlotRole.FOOTER})
        return slot

    return [
        pattern.model_copy(update={"slots": [footer(slot) for slot in pattern.slots]})
        for pattern in patterns
    ]


def _sequence_ordinals(pattern: Pattern, tree) -> Pattern:
    """Номера по порядку в фигурах одного размера — порядковые номера, не показатели.

    «01», «02», «03» в кружках шаблона экзаменов разбирались местом под
    показатель (короткий крупный текст): туда садились «47,3 млн руб.» и
    пункты списка, и текст шёл колонкой шириной в дюйм. «02» стоит со
    сдвигом, поэтому «01» и «03» попадали в повторитель, а «02» — в
    отдельное место: ряд виден только по всем фигурам слайда сразу. Число в
    ряду 1, 2, 3… с соседями того же размера — номер элемента. Одиночное «7»
    с подписью (`vk_tech`) остаётся показателем.
    """
    numbers: list[tuple[Box, int]] = []
    for element, box, _ in iter_shapes(tree):
        body = element.find(f"{{{P_NS}}}txBody")
        if box is None or body is None:
            continue
        text = "".join(node.text or "" for node in body.iter(f"{{{A_NS}}}t")).strip()
        if text.isdigit() and len(text) <= 2:
            numbers.append((box, int(text)))
    ordinal_boxes: list[Box] = []
    for box, _ in numbers:
        twins = [
            (other, value)
            for other, value in numbers
            if abs(other.w - box.w) <= box.w // 10 and abs(other.h - box.h) <= box.h // 10
        ]
        values = sorted(value for _, value in twins)
        if len(twins) >= 2 and values == list(range(1, len(twins) + 1)):
            ordinal_boxes.append(box)
    if not ordinal_boxes:
        return pattern

    def on_ordinal(box: Box) -> bool:
        return any(_covers_most(box, other) and _covers_most(other, box) for other in ordinal_boxes)

    def marked(slot: Slot, offsets: list[tuple[int, int]]) -> Slot:
        if slot.role in (SlotRole.TITLE, SlotRole.SLIDE_NUMBER, SlotRole.ORDINAL):
            return slot
        if any(
            on_ordinal(slot.box.model_copy(update={"x": slot.box.x + dx, "y": slot.box.y + dy}))
            for dx, dy in offsets
        ):
            return slot.model_copy(update={"role": SlotRole.ORDINAL})
        return slot

    return pattern.model_copy(
        update={
            "slots": [marked(slot, [(0, 0)]) for slot in pattern.slots],
            "repeaters": [
                repeater.model_copy(
                    update={
                        "item_slots": [
                            marked(slot, repeater.member_offsets or [(0, 0)])
                            for slot in repeater.item_slots
                        ]
                    }
                )
                for repeater in pattern.repeaters
            ],
        }
    )


def _figure_pictures(tree) -> list[Box]:
    """Картинки слайда, на которых стоит число-показатель шаблона."""
    shapes = [(element, box) for element, box, _ in iter_shapes(tree) if box is not None]
    def own_text(element) -> str:
        # Только свой `txBody`: полный обход захватывает и запасные ветки
        # `AlternateContent`, и «10%» читался как «10%10%10%».
        body = element.find(f"{{{P_NS}}}txBody")
        if body is None:
            return ""
        return "".join(node.text or "" for node in body.iter(f"{{{A_NS}}}t"))

    numbers = [
        box
        for element, box in shapes
        if etree.QName(element).localname == "sp" and is_figure_text(own_text(element))
    ]
    pictures = [box for element, box in shapes if etree.QName(element).localname == "pic"]
    figures = [box for box in pictures if any(_covers_most(n, box) for n in numbers)]
    # И части той же фигуры: дуга кольца — отдельная картинка внутри него.
    return [box for box in pictures if any(_covers_most(box, f) for f in figures)]


_DECOR_TAGS = frozenset({"sp", "pic", "cxnSp"})
_THIN = 91_440  # 0.1″: линия
_ICON = 548_640  # 0.6″: значок


# Доля слайда, начиная с которой `_decor` перестаёт разбирать фигуру сам:
# по габаритам крупной фигуры не отличить наложение от замысла. Ровно эти
# фигуры и приходят разобранными по содержанию из `parse.surfaces`.
DECOR_AREA_SHARE = 0.5


# Заливка мельче этой доли слайда — не фон текста, а точка оформления.
_GROUND_MIN_AREA = 0.002


def _background_sources(slide):
    """Где искать фон-картинку: слайд, его макет, образец — в этом порядке."""
    layout = slide.slide_layout
    master = layout.slide_master
    return [
        (tokens_mod.own_background(slide._element), slide.part),
        (tokens_mod.own_background(layout._element), layout.part),
        (tokens_mod.own_background(master._element), master.part),
    ]


def _background_decor(slide, slide_w: int, slide_h: int) -> list[Box]:
    """Графика внутри фона-картинки: её тоже нельзя закрывать текстом."""
    for background, part in _background_sources(slide):
        found = busy_regions(background, part, slide_w, slide_h)
        if found:
            return found
    return []


def _grounds(
    tree, theme, clr_map, base, slide_w: int, slide_h: int, backgrounds=()
) -> list[Ground]:
    """Всё залитое на слайде-доноре: на чём угодно из этого может лечь текст.

    Записывается фактом разбора, потому что цвет текста нельзя решать по
    слоту: у блока слота может не быть (свободная полоса, элемент
    повторителя, подпись под числом), а фон под ним есть всегда. Вёрстка
    выбирает цвет по итоговой рамке и по этим областям.
    """
    slide = Box(x=0, y=0, w=slide_w, h=slide_h)
    # Фон-картинка лежит ниже всего: любая заливка её закрывает.
    cells: list[tuple[Box, str]] = []
    for background, part in backgrounds:
        cells = background_grid(background, part, slide_w, slide_h)
        if cells:
            break
    found = [Ground(box=box, colors=[Color(rgb=rgb)], z=-1) for box, rgb in cells]
    return found + [
        Ground(box=box, colors=colors, z=order)
        for box, colors, order in tokens_mod.ground_regions(
            tree, slide, theme, clr_map, base, min_share=_GROUND_MIN_AREA
        )
    ]


def _slide_obstacles(tree, part, pattern: Pattern, slide_w: int, slide_h: int) -> list[Box]:
    """Крупные элементы самого слайда-образца, поверх которых писать нельзя.

    `_decor` отбирает по площади и намеренно грубо: фигура крупнее половины
    слайда чаще всего заливка, по которой текст и должен лежать. Но площадь
    не отличает заливку от фотографии, поэтому крупное приходит сюда — уже
    расклассифицированным по содержанию (`parse.surfaces`).

    Берётся ровно то, что `_decor` отбросил, — крупнее порога. Всё, что мельче,
    он разбирает сам и знает про слоты композиции; дублировать его работу
    здесь значит объявить препятствием карточки самой композиции.

    Места композиции исключаются и тут: слот — это место под наш текст, а не
    препятствие для него. Место под фотографию — тоже: донорское фото туда
    не переносится, там будет наша картинка.
    """
    visible = [
        surface
        for surface in surfaces.collect_tree(tree, part, slide_w, slide_h)
        if surface.box.area >= DECOR_AREA_SHARE * slide_w * slide_h
    ]
    keep = [slot.box for slot in pattern.slots] + list(pattern.photo_slots)
    keep += [
        Box(x=slot.box.x + dx, y=slot.box.y + dy, w=slot.box.w, h=slot.box.h)
        for repeater in pattern.repeaters
        for dx, dy in repeater.member_offsets
        for slot in repeater.item_slots
    ]
    # Сначала отсев, потом защита. Места композиции — не препятствие, и
    # узнать их можно только по исходной рамке: защита её растит и
    # склеивает соседние, и карточки самого слайда переставали совпадать
    # со своими слотами. На шаблоне экзаменов из-за этого три карточки
    # оказались запрещены, и три пункта уехали в одну.
    own = [
        surface
        for surface in visible
        if not any(
            _same_rect(surface.box, slot) or _covers_most(surface.box, slot)
            for slot in keep
        )
    ]
    return surfaces.protected(own, slide_w, slide_h)


def _decor(tree, pattern, slide_w: int, slide_h: int) -> list[Box]:
    """Графика донора, на которую текст не должен заходить.

    Фигуры без своего текста: линии, иконки, картинки. Не входят подложки —
    фигуры, под которыми целиком лежит текстовое место (карточка, панель), —
    и фон во весь слайд: писать на них текст и задумано.
    """
    texts = [slot.box for slot in pattern.slots if slot.role not in _NOT_TEXT_ROLES]
    for repeater in pattern.repeaters:
        for dx, dy in repeater.member_offsets:
            texts += [
                Box(x=slot.box.x + dx, y=slot.box.y + dy, w=slot.box.w, h=slot.box.h)
                for slot in repeater.item_slots
                if slot.role not in _NOT_TEXT_ROLES
            ]
    found = []
    for element, box, _ in iter_shapes(tree):
        if box is None or etree.QName(element).localname not in _DECOR_TAGS:
            continue
        body = element.find(f"{{{P_NS}}}txBody")
        if body is not None and "".join(n.text or "" for n in body.iter(f"{{{A_NS}}}t")).strip():
            continue
        if box.w * box.h >= DECOR_AREA_SHARE * slide_w * slide_h:
            continue
        if any(_covers_most(text, box) for text in texts):
            continue
        # Только линии и значки. Крупная фигура — поверхность, на которой
        # донор и сам ставит текст (кольцо вокруг заголовка `vk_education`,
        # панель под заголовком `vk_tech`); по её габаритам не отличить
        # наложение от замысла.
        if min(box.w, box.h) > _THIN and max(box.w, box.h) > _ICON:
            continue
        found.append(box)
    return found


# Линия для записи: тонкая горизонталь под подписью («____» шаблона-бланка).
_WRITING_LINE = 27_432  # 0.03″
# Роли, чья рамка растёт вниз по свободному месту.
_GROWING_ROLES = frozenset({SlotRole.BODY, SlotRole.CAPTION, SlotRole.KPI_LABEL})
# Зазор между выросшей рамкой и тем, во что она упёрлась.
_GROWTH_GAP = 45_720  # 0.05″
# Ниже этой доли слайда — служебная зона: номер страницы, колонтитул.
_GROWTH_FLOOR = 0.9


def _is_writing_line(box: Box) -> bool:
    return box.h <= _WRITING_LINE and box.w > 10 * max(1, box.h)


# Доля рамки, перекрытая кругом, при которой надпись — метка узла.
_LABEL_ON_NODE = 0.3


def _shared_area(a: Box, b: Box) -> int:
    width = min(a.right, b.right) - max(a.x, b.x)
    height = min(a.bottom, b.bottom) - max(a.y, b.y)
    return max(0, width) * max(0, height)


def _same_rect(a: Box, b: Box) -> bool:
    return (a.x, a.y, a.w, a.h) == (b.x, b.y, b.w, b.h)


def _is_line_shape(element) -> bool:
    """Соединитель или линия: габарит диагонали — не препятствие для текста."""
    if etree.QName(element).localname == "cxnSp":
        return True
    geometry = element.find(f".//{{{A_NS}}}prstGeom")
    preset = (geometry.get("prst") or "").lower() if geometry is not None else ""
    return "line" in preset or "connector" in preset


def _room_below(box: Box, shapes: list[tuple[Box, bool]], slide_h: int) -> int:
    """Где кончается свободное место под рамкой: высота, до которой ей расти.

    Останавливает любая фигура ниже рамки, пересекающая её по горизонтали,
    кроме линий для записи: писать поверх них шаблон и задумал. Подложка,
    внутри которой стоит рамка, останавливает своим низом.
    """
    limit = int(slide_h * _GROWTH_FLOOR)
    for other, round_ in shapes:
        if _same_rect(other, box) or _is_writing_line(other):
            continue
        across = min(other.right, box.right) - max(other.x, box.x)
        if across <= 0:
            continue
        # Надпись на круге — метка узла («Проблема» в кружке схемы): круг её
        # не растянет, а выросшая рамка забрала бы место описания под ним.
        if round_ and _shared_area(box, other) >= box.area * _LABEL_ON_NODE:
            return box.h
        if _covers_most(box, other):
            limit = min(limit, other.bottom)
        elif other.y >= box.bottom - _GROWTH_GAP:
            limit = min(limit, other.y)
    return max(box.h, limit - box.y - _GROWTH_GAP)


def _grown_frames(
    tree,
    pattern: Pattern,
    slide_w: int,
    slide_h: int,
    obstacles: list[Box] | None = None,
) -> Pattern:
    """Рамки текста шаблона-бланка — во всё свободное место под ними.

    Шаблон без плейсхолдеров набирает каждую надпись в рамку ровно под
    своё слово: «Что нужно изменить?» — 0.31″, под ней линия для записи.
    Предложение туда не помещается ни на каком кегле, и вёрстка уходила на
    обложку и финал — единственные высокие рамки. Место под текст в таком
    шаблоне — это рамка вместе с пустым пространством и линиями под ней.
    Рамка элемента повторителя растёт одинаково у всех элементов — на
    наименьший из их запасов.
    """
    shapes = [
        (box, element.find(f".//{{{A_NS}}}prstGeom[@prst='ellipse']") is not None)
        for element, box, _ in iter_shapes(tree)
        if box is not None
        and box.area < 0.5 * slide_w * slide_h
        and not _is_line_shape(element)
    ]
    # Оформление макета останавливает рост так же, как фигуры слайда. Без
    # этого рамка подзаголовка обложки `vk_tech` вырастала с 0.31 до 0.94
    # дюйма прямо поверх картинки макета: препятствие лежало в макете, а
    # `_room_below` смотрел только на слайд. Порог площади здесь не нужен —
    # эти рамки уже отобраны как запрещающие текст.
    shapes += [(box, False) for box in (obstacles or [])]

    def grow(slot: Slot, offsets: list[tuple[int, int]]) -> Slot:
        if slot.role not in _GROWING_ROLES:
            return slot
        room = min(
            _room_below(
                slot.box.model_copy(update={"x": slot.box.x + dx, "y": slot.box.y + dy}),
                shapes,
                slide_h,
            )
            for dx, dy in offsets
        )
        if room <= slot.box.h:
            return slot
        return slot.model_copy(update={"box": slot.box.model_copy(update={"h": room})})

    texts = {
        (box.x, box.y, box.w, box.h): "".join(
            node.text or "" for node in element.iter(f"{{{A_NS}}}t")
        )
        for element, box, _ in iter_shapes(tree)
        if box is not None
    }

    def ordinal(slot: Slot, offsets: list[tuple[int, int]]) -> Slot:
        """Номер элемента: у элементов донора там стоят 1, 2, 3… по порядку."""
        found = [
            texts.get((slot.box.x + dx, slot.box.y + dy, slot.box.w, slot.box.h), "").strip()
            for dx, dy in offsets
        ]
        if len(found) < 2 or not all(text.isdigit() for text in found):
            return slot
        if [int(text) for text in found] != list(range(1, len(found) + 1)):
            return slot
        return slot.model_copy(update={"role": SlotRole.ORDINAL})

    grown_slots = [grow(slot, [(0, 0)]) for slot in pattern.slots]
    repeaters = []
    for repeater in pattern.repeaters:
        offsets = repeater.member_offsets or [(0, 0)]
        items = [grow(ordinal(slot, offsets), offsets) for slot in repeater.item_slots]
        # Рамка элемента растёт вместе с его текстом: иначе рендер не узнавал
        # в выросшем тексте содержание элемента и убирал элемент целиком.
        item_box = _union_all([repeater.item_box, *(slot.box for slot in items)])
        frames = [
            _union_all(
                [
                    frame,
                    *(
                        slot.box.model_copy(update={"x": slot.box.x + dx, "y": slot.box.y + dy})
                        for slot in items
                    ),
                ]
            )
            for frame, (dx, dy) in zip(
                repeater.member_frames, repeater.member_offsets, strict=False
            )
        ]
        repeaters.append(
            repeater.model_copy(
                update={"item_slots": items, "item_box": item_box, "member_frames": frames}
            )
        )
    grown = [slot.box for slot in grown_slots] + [
        slot.box.model_copy(update={"x": slot.box.x + dx, "y": slot.box.y + dy})
        for repeater in repeaters
        for dx, dy in (repeater.member_offsets or [(0, 0)])
        for slot in repeater.item_slots
    ]
    # Линии для записи под выросшей рамкой — не декор, который текст обходит:
    # текст ляжет на их место, а рендер их уберёт.
    decor = [
        box
        for box in pattern.decor
        if not (_is_writing_line(box) and any(_covers_most(box, frame) for frame in grown))
    ]
    return pattern.model_copy(
        update={"slots": grown_slots, "repeaters": repeaters, "decor": decor}
    )


def _as_bookend(pattern: Pattern, kind: PatternClass, slide_h: int) -> Pattern:
    """Обложка или финал шаблона-бланка: надписи над заголовком — не места.

    Как у обычной обложки (`bookend_pattern`): «РЕДАКТИРУЕМЫЙ ШАБЛОН» над
    заголовком и плашка «НАЗВАНИЕ ПРОЕКТА» в углу — оформление шаблона.
    Подзаголовок титула уходил в строку над заголовком и налезал на него.
    Так же — надписи рядом с заголовком (плашка начинается на его высоте) и
    в нижней служебной полосе («ДОРОЖНАЯ КАРТА ПРОЕКТА» у номера страницы).
    """
    title = next((slot for slot in pattern.slots if slot.role is SlotRole.TITLE), None)
    if title is None:
        return pattern.model_copy(update={"pattern_class": kind})
    slots = [
        slot.model_copy(update={"role": SlotRole.DECOR})
        if slot is not title
        and (slot.box.y < title.box.bottom or slot.box.y >= slide_h * _GROWTH_FLOOR)
        else slot
        for slot in pattern.slots
    ]
    return pattern.model_copy(update={"pattern_class": kind, "slots": slots})


# Ячейки одной строки стоят на одной высоте с таким допуском.
_ROW_TOLERANCE_EMU = 91_440  # 0.1″
# Шапка — не дальше этого над первой строкой.
_HEADER_REACH_EMU = 1_097_280  # 1.2″
_CELL_ROLES = frozenset(
    {SlotRole.BODY, SlotRole.CAPTION, SlotRole.KPI_LABEL, SlotRole.KPI_VALUE}
)


def _overlap_x(a: Box, b: Box) -> int:
    return max(0, min(a.right, b.right) - max(a.x, b.x))


def _table_grid(pattern: Pattern) -> Pattern:
    """Таблица из прямоугольников: строки-повторитель и шапка над ними.

    Бланк рисует таблицу фигурами — ячейки шапки залиты, строки одинаковые, в
    строке надписи по колонкам. Нативной таблицы нет, и без этого разбора
    шапка была пятью местами под текст: список ложился в ячейку «№» шириной
    0.69″. Ячейка шапки — надпись над первой строкой, по горизонтали на
    колонке ячейки строки.
    """
    for repeater in pattern.repeaters:
        if repeater.axis != "vertical":
            continue
        texts = [slot for slot in repeater.item_slots if slot.role in _CELL_ROLES]
        if len(texts) < 2:
            continue
        top = min(slot.box.y for slot in texts)
        row = sorted(
            (slot for slot in texts if slot.box.y - top <= _ROW_TOLERANCE_EMU),
            key=lambda slot: slot.box.x,
        )
        if len(row) < 2:
            continue
        first = repeater.member_frames[0].y if repeater.member_frames else repeater.item_box.y
        heads = [
            slot
            for slot in pattern.slots
            if slot.role in _CELL_ROLES
            and first - _HEADER_REACH_EMU <= slot.box.y
            and slot.box.y + min(slot.box.h, _ROW_TOLERANCE_EMU * 5) <= first + _ROW_TOLERANCE_EMU
        ]
        mapped = []
        for cell in row:
            best = max(heads, key=lambda head: _overlap_x(head.box, cell.box), default=None)
            if best is None or not _overlap_x(best.box, cell.box) or best in mapped:
                break
            mapped.append(best)
        if len(mapped) != len(row):
            continue
        number = None
        ordinals = [slot for slot in repeater.item_slots if slot.role is SlotRole.ORDINAL]
        if ordinals:
            number = max(
                (head for head in heads if head not in mapped),
                key=lambda head: _overlap_x(head.box, ordinals[0].box),
                default=None,
            )
            if number is not None and not _overlap_x(number.box, ordinals[0].box):
                number = None
        # Шапка — больше не место под пункты: её заполняет только таблица.
        header_ids = {slot.id for slot in mapped} | ({number.id} if number else set())
        slots = [
            slot.model_copy(update={"role": SlotRole.UNKNOWN}) if slot.id in header_ids else slot
            for slot in pattern.slots
        ]
        return pattern.model_copy(
            update={
                "slots": slots,
                "table_grid": TableGrid(
                    repeater_id=repeater.id,
                    header_slot_ids=[slot.id for slot in mapped],
                    cell_slot_ids=[slot.id for slot in row],
                    number_header_id=number.id if number else None,
                ),
            }
        )
    return pattern


def _union_all(boxes: list[Box]) -> Box:
    left, top = min(b.x for b in boxes), min(b.y for b in boxes)
    right, bottom = max(b.right for b in boxes), max(b.bottom for b in boxes)
    return Box(x=left, y=top, w=right - left, h=bottom - top)


def _covers_most(inner: Box, outer: Box) -> bool:
    width = min(inner.right, outer.right) - max(inner.x, outer.x)
    height = min(inner.bottom, outer.bottom) - max(inner.y, outer.y)
    return width > 0 and height > 0 and width * height >= 0.9 * inner.area


def _overlap(a: Box, b: Box) -> bool:
    return min(a.right, b.right) > max(a.x, b.x) and min(a.bottom, b.bottom) > max(a.y, b.y)


def _text_of_member(repeater, dx: int, dy: int) -> Box:
    """Главная текстовая рамка элемента повторителя с этим сдвигом.

    Самая крупная, а не объединение всех: номер в кружке над карточкой
    выходит за её край, и по объединению карточка не находилась подложкой.
    """
    main = max(
        (slot.box for slot in repeater.item_slots),
        key=lambda box: box.area,
        default=repeater.item_box,
    )
    return Box(x=main.x + dx, y=main.y + dy, w=main.w, h=main.h)


def _member_align(tree, box: Box):
    """Выравнивание текста фигуры донора, стоящей в этой рамке."""
    tolerance = 45_720  # 1/20 дюйма
    for element, shape_box, _ in iter_shapes(tree):
        if shape_box is None:
            continue
        if all(
            abs(a - b) <= tolerance
            for a, b in zip(
                (shape_box.x, shape_box.y, shape_box.w, shape_box.h),
                (box.x, box.y, box.w, box.h),
                strict=True,
            )
        ):
            align = text_align(element)
            if align is not None:
                return align
    return None


def _list_indent(prs: PresentationObject) -> int:
    """Висячий отступ маркера списка из спискового стиля образца, EMU.

    Первый уровень `bodyStyle` — тот, которым шаблон набирает списки на
    содержательных слайдах. Отступ наследуется молча: на фигуре донора его
    нет, а в рендере он есть.
    """
    found = 0
    for master in prs.slide_masters:
        node = master._element.find(
            f".//{{{P_NS}}}txStyles/{{{P_NS}}}bodyStyle/{{{A_NS}}}lvl1pPr"
        )
        if node is None:
            continue
        try:
            found = max(found, int(node.get("marL", "0")))
        except ValueError:
            continue
    return found


def _parser_fingerprint() -> str:
    """Отпечаток исходников разбора: правка парсера сбрасывает кэш сама.

    Раньше ключом кэша был только хэш файла шаблона, и исправленный разбор
    молча не применялся к уже разобранным шаблонам — ни у разработчика, ни в
    CI с сохранённым кэшем. Ручной номер версии забывают поднять; отпечаток
    исходников не забывает.
    """
    digest = hashlib.sha256()
    package = Path(__file__).resolve().parent
    for source in sorted(package.glob("*.py")):
        digest.update(source.read_bytes())
    # Схема входит в отпечаток наравне с разбором: `TemplateSpec` пополнился
    # ограничениями макета, и кэш, записанный до этого, описывает шаблон без
    # них. Без этой строки старый кэш продолжал бы отдавать композиции, для
    # которых слайд — пустой холст.
    digest.update((package.parent / "schemas" / "template_spec.py").read_bytes())
    # И общие типы: `TextStyle` и `Box` лежат внутри `TemplateSpec`, и их
    # смысл — часть того, что записано в кэш.
    digest.update((package.parent / "schemas" / "common.py").read_bytes())
    return digest.hexdigest()[:12]


def parse_template(
    path: str | Path,
    cache_dir: str | Path | None = None,
    font_dir: str | Path | None = None,
) -> TemplateSpec:
    """Разбирает `.pptx`. Повторный разбор того же файла берётся из кэша."""
    path = Path(path)
    cache_path: Path | None = None

    if cache_dir is not None:
        cache_path = Path(cache_dir) / f"{file_sha256(path)}-{_parser_fingerprint()}.json"
        if cache_path.exists():
            try:
                return TemplateSpec.model_validate_json(cache_path.read_text("utf-8"))
            except (ValueError, json.JSONDecodeError):
                # Кэш от прежней версии схемы: разбираем заново и перезаписываем.
                cache_path.unlink(missing_ok=True)

    spec = _parse(path, Path(font_dir) if font_dir else None)

    if cache_path is not None:
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        cache_path.write_text(spec.model_dump_json(), encoding="utf-8")
    return spec
