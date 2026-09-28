"""Плотность, целостность и числа — то, что проверяется без модели.

**Числа проверяются детерминированно, и это принципиально.** Спросить модель
«сходится ли 4.6 раза» значит получить мнение; пересчитать по формуле, которую
модель же и предъявила, значит получить ответ. Цитируемое число сверяется с
фактом контент-пакета, производное — пересчитывается.
"""

from __future__ import annotations

import re
from pathlib import Path

from deckwright.audit.registry import check
from deckwright.plan.figures import verify as verify_figure
from deckwright.schemas import (
    DeckIR,
    DeckPlan,
    FixKind,
    Issue,
    PatternClass,
    ProposedFix,
    SlideIR,
    SlotRole,
)

# Доля площади слайда, ниже которой он считается пустоватым. Порог
# информационный: воздушный вариант имеет право быть просторным.
MIN_FILL_RATIO = 0.25


def _issue(check_id: str, slide_index: int, message: str, **kwargs) -> Issue:
    spec = check(check_id)
    fix = kwargs.pop("fix", None)
    return Issue(
        check_id=check_id,
        kind=spec.kind,
        category=spec.category,
        severity=spec.severity,
        slide_index=slide_index,
        message=message,
        fix=fix or ProposedFix(kind=FixKind.NONE),
        **kwargs,
    )


def density(
    slide: SlideIR, max_bullets: int, max_words: int, min_fill: float = MIN_FILL_RATIO
) -> list[Issue]:
    """Пороги плотности из Приложения 1: сколько пунктов и какой длины."""
    found: list[Issue] = []
    for element in slide.all_elements():
        if element.text is None:
            continue
        bullets = [p for p in element.text.paragraphs if p.bullet]
        if len(bullets) > max_bullets:
            found.append(
                _issue(
                    "density.too_many_bullets",
                    slide.index,
                    f"{element.id}: пунктов {len(bullets)} при пороге {max_bullets}",
                    element_ids=[element.id],
                    bbox=element.box,
                    fix=ProposedFix(
                        kind=FixKind.ASSISTED,
                        description="разнести часть пунктов на другой слайд",
                        action="split_bullets",
                        params={"element_id": element.id},
                    ),
                )
            )
        for index, paragraph in enumerate(element.text.paragraphs):
            words = len(paragraph.text.split())
            if words <= max_words:
                continue
            found.append(
                _issue(
                    "density.bullet_too_long",
                    slide.index,
                    f"{element.id}, абзац {index + 1}: {words} слов при пороге {max_words}",
                    element_ids=[element.id],
                    bbox=element.box,
                    fix=ProposedFix(
                        kind=FixKind.ASSISTED,
                        description="сократить формулировку",
                        action="shorten_paragraph",
                        params={"element_id": element.id, "paragraph": index},
                    ),
                )
            )
    return found


def title_only(deck: DeckIR, spec=None) -> list[Issue]:
    """Слайд содержания, на котором кроме заголовка ничего нет.

    Заполненность (`fill_ratio`) считает все фигуры слайда, и логотип,
    линия и картинка донора её набирают: шесть слайдов питча Fibonacci
    из одних заголовков («Рынок растёт», «LTV в 4 раза выше CAC») прошли
    с находкой уровня info. Здесь — по `SlideIR`: есть ли у слайда хоть
    один элемент с текстом, график или таблица, кроме заголовка и
    подзаголовка. Обложка и финал — не содержание, их не проверяем.
    """
    found: list[Issue] = []
    bookends = spec.bookend_ids if spec is not None else set()
    # Раздел — заголовок по замыслу: разделитель без содержания не потеря.
    if spec is not None:
        bookends |= {
            pattern.id
            for pattern in spec.patterns
            if pattern.pattern_class in (PatternClass.TITLE, PatternClass.SECTION)
        }
    last = len(deck.slides)
    for slide in deck.slides:
        if slide.pattern_id in bookends or (not bookends and slide.index in (1, last)):
            continue
        content = [
            element
            for element in slide.all_elements()
            if element.role not in _HEADER_ROLES
            and (
                element.chart is not None
                or element.table is not None
                or (
                    element.text is not None
                    and any(p.text.strip() for p in element.text.paragraphs)
                )
            )
        ]
        if not content:
            found.append(
                _issue(
                    "density.title_only",
                    slide.index,
                    "на слайде только заголовок: содержание слайда потерялось",
                )
            )
    return found


_HEADER_ROLES = frozenset(
    {SlotRole.TITLE, SlotRole.SUBTITLE, SlotRole.FOOTER, SlotRole.SLIDE_NUMBER}
)


def fill_ratio(
    pptx_path: str | Path, deck: DeckIR, min_fill: float = MIN_FILL_RATIO
) -> list[Issue]:
    """Слишком пустой слайд: содержание где-то потерялось.

    Считается по **собранному `.pptx`**, а не по `SlideIR`. После того как
    композиция стала клонироваться с донора, представление вёрстки описывает
    лишь малую часть слайда — карточки, картинки и декор живут в пакете. Счёт
    по IR давал «пусто» на десяти слайдах из двенадцати, при том что слайды
    заполнены. Мерить надо то, что видит человек.
    """
    from pptx import Presentation

    area = deck.slide_width_emu * deck.slide_height_emu
    if area <= 0:
        return []

    found: list[Issue] = []
    presentation = Presentation(str(pptx_path))
    for index, slide in enumerate(presentation.slides, start=1):
        used = sum(
            (shape.width or 0) * (shape.height or 0)
            for shape in slide.shapes
            if shape.left is not None
        )
        ratio = used / area
        if ratio >= min_fill:
            continue
        found.append(
            _issue(
                "density.slide_too_empty",
                index,
                f"слайд заполнен на {100 * ratio:.0f}% при пороге {100 * min_fill:.0f}%",
            )
        )
    return found


def duplicate_slides(deck: DeckIR) -> list[Issue]:
    """Два слайда с одинаковым текстом — потерянная правка, а не повтор.

    Сравнивается текст, а не композиция: один и тот же тезис, свёрстанный
    дважды по-разному, всё равно повтор.
    """

    def fingerprint(slide: SlideIR) -> str:
        parts = [
            paragraph.text.strip().lower()
            for element in slide.all_elements()
            if element.text is not None
            for paragraph in element.text.paragraphs
            if paragraph.text.strip()
        ]
        return "|".join(sorted(parts))

    seen: dict[str, int] = {}
    found: list[Issue] = []
    for slide in deck.slides:
        key = fingerprint(slide)
        if not key:
            continue
        if key in seen:
            found.append(
                _issue(
                    "integrity.duplicate_slides",
                    slide.index,
                    f"слайд повторяет слайд {seen[key]} слово в слово",
                )
            )
            continue
        seen[key] = slide.index
    return found


def package(pptx_path: str | Path) -> list[Issue]:
    """Целостность пакета и слайды-картинки — по собранному файлу, не по IR."""
    from deckwright.render.package_check import check_package
    from deckwright.render.pptx_writer import slide_is_single_image

    found: list[Issue] = []
    report = check_package(pptx_path)
    if not report.ok:
        found.append(
            _issue(
                "integrity.package_broken",
                1,
                "пакет повреждён: " + "; ".join(report.problems[:3]),
            )
        )
    for index in slide_is_single_image(pptx_path):
        found.append(
            _issue(
                "integrity.slide_is_single_image",
                index,
                "слайд состоит из одной картинки — ТЗ такое не засчитывает",
            )
        )
    return found


# Имена, которые Office даёт рядам и категориям графика по умолчанию. Это
# данные шаблона-донора, а не ответ модели.
_DEFAULT_CHART_NAME = re.compile(r"^(Ряд|Серия|Series|Категория|Category)\s*\d+$", re.I)


def donor_data(pptx_path: str | Path, deck: DeckIR, spec=None) -> list[Issue]:
    """Данные донора, доехавшие до готовой колоды: ложные цифры.

    Рыбный график «Ряд 1/2/3» рядом с нашим и кольцо «10%» — визуализация
    данных, которых в контент-пакете нет. Детерминированные проверки по IR
    их не видят: в IR этих фигур нет, они приезжают клонированием донора.
    Поэтому проверяется собранный файл: у графиков — имена рядов и категорий,
    у текста — числа-показатели, которых нет в нашем содержании слайда.
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    from deckwright.parse.patterns import is_figure_text
    from deckwright.render.pptx_writer import _same_box, _shape_box_of, _similar_size

    found: list[Issue] = []
    patterns = {pattern.id: pattern for pattern in getattr(spec, "patterns", [])}
    presentation = Presentation(str(pptx_path))
    for slide_ir, slide in zip(deck.slides, presentation.slides, strict=False):
        pattern = patterns.get(slide_ir.pattern_id)
        data_boxes = [
            element.box
            for element in slide_ir.all_elements()
            if element.chart is not None or element.table is not None
        ]
        for shape in _walk(slide.shapes):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            box = _shape_box_of(shape)
            # Картинка донора с его числом: кольцо «10%» без числа — всё ещё
            # чужая доля. И картинка того же размера, что наш график рядом, —
            # рыбный график дизайнера.
            if pattern is not None and any(
                _same_box(box, figure) for figure in pattern.figure_pictures
            ):
                found.append(
                    _issue(
                        "integrity.donor_data_leftover",
                        slide_ir.index,
                        "картинка донора с его числом: визуализация данных, которых нет",
                    )
                )
            elif any(_similar_size(box, data) for data in data_boxes):
                found.append(
                    _issue(
                        "integrity.donor_data_leftover",
                        slide_ir.index,
                        "картинка донора размером с наш график: рыбный график рядом с нашим",
                    )
                )
        ours = {
            paragraph.text.strip()
            for element in slide_ir.all_elements()
            if element.text is not None
            for paragraph in element.text.paragraphs
        }
        for shape in _walk(slide.shapes):
            if getattr(shape, "has_chart", False) and shape.has_chart:
                plot = shape.chart.plots[0] if shape.chart.plots else None
                names = [series.name for series in plot.series] if plot else []
                names += list(plot.categories) if plot else []
                fish = [name for name in names if _DEFAULT_CHART_NAME.match(str(name))]
                if fish:
                    found.append(
                        _issue(
                            "integrity.donor_data_leftover",
                            slide_ir.index,
                            f"график донора с рыбными данными: {', '.join(map(str, fish[:3]))}",
                        )
                    )
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if is_figure_text(text) and text not in ours:
                    found.append(
                        _issue(
                            "integrity.donor_data_leftover",
                            slide_ir.index,
                            f"число донора {text!r}: в содержании слайда его нет",
                        )
                    )
    return found


def _walk(shapes):
    """Фигуры слайда вместе с содержимым групп."""
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):
            yield from _walk(shape.shapes)


def figures(plan: DeckPlan, pack) -> list[Issue]:
    """Числа: цитата сверяется с фактом, производное пересчитывается формулой.

    Обе проверки детерминированные и потому пригодны для автоправки в отличие
    от вопроса модели «похоже ли это на правду».
    """
    found: list[Issue] = []
    for slide in plan.slides:
        for figure in slide.figures:
            problem = verify_figure(figure, pack)
            if problem is None:
                continue
            check_id = (
                "content.derived_figure_wrong"
                if figure.formula
                else "content.figure_not_in_sources"
            )
            found.append(
                _issue(
                    check_id,
                    slide.index,
                    f"{figure.text!r}: {problem}",
                )
            )
    return found


# Целые до этого значения — счёт и нумерация («3 шага», «Этап 2»), а не факт.
SMALL_COUNT = 10


def undeclared_numbers(plan: DeckPlan, pack) -> list[Issue]:
    """Число в тексте слайда, которого во входе нет: ни в фактах, ни в рядах.

    `figure_not_in_sources` проверяет только числа, которые модель сама
    объявила в `figures`. Число, написанное в пункте и не объявленное,
    проходило молча — а это ровно выдуманная цифра. Здесь проверяется весь
    текст слайда: заголовок, подзаголовок, пункты, ячейки таблиц. Число
    засчитывается, если оно есть во входе (с точностью до масштаба записи)
    или объявлено выведенным и сходится со своей формулой.
    """
    from deckwright.content.grounding import source_numbers, ungrounded

    sources = [pack.brief.topic, pack.brief.request, pack.brief.goal, pack.brief.audience]
    sources += [pack.brief.extra_instructions]
    for fact in pack.facts:
        sources.append(fact.text)
        if fact.value is not None:
            sources.append(f"{fact.value:g}")
    for series in pack.series:
        sources.append(series.name)
        sources += [f"{point.label} {point.value:g}" for point in series.points]
    sources += [quote.text for quote in pack.quotes]
    grounded = source_numbers("\n".join(sources))

    found: list[Issue] = []
    for slide in plan.slides:
        allowed = set(grounded)
        for figure in slide.figures:
            if verify_figure(figure, pack) is None:
                allowed |= source_numbers(figure.text)
        texts = [slide.takeaway_title, slide.subtitle]
        for block in slide.blocks:
            texts += [block.heading, *block.items]
            if block.table is not None:
                texts += list(block.table.columns)
                texts += [cell for row in block.table.rows for cell in row]
        invented = sorted({n for text in texts for n in ungrounded(text, allowed, SMALL_COUNT)})
        if invented:
            found.append(
                _issue(
                    "content.undeclared_number",
                    slide.index,
                    "числа "
                    + ", ".join(f"{n:g}" for n in invented)
                    + " на слайде, а во входе их нет",
                )
            )
    return found


def donor_photos(pptx_path: str | Path, deck: DeckIR, spec=None) -> list[Issue]:
    """Фото на собранном слайде — всегда фото донора.

    Колода фото не создаёт (генерации картинок нет), поэтому любое фото в
    ней приехало клонированием композиции: студент за компьютером на слайде
    о продажах кофеен. Контекстный вопрос «картинки по теме» этого не поймал
    — модель видела уместную по виду иллюстрацию, темы колоды не знала и по
    правилу «сомнение — это да» отвечала «да». Здесь — детерминированно, по
    самому изображению (`parse.pictures`), а не по имени фигуры.
    """
    from pptx import Presentation

    from deckwright.parse.pictures import photo_boxes

    found: list[Issue] = []
    presentation = Presentation(str(pptx_path))
    width, height = presentation.slide_width, presentation.slide_height
    for slide_ir, slide in zip(deck.slides, presentation.slides, strict=False):
        # На обложке и финале иллюстрации — оформление шаблона.
        bookend = spec is not None and slide_ir.pattern_id in spec.bookend_ids
        photos = photo_boxes(
            slide.shapes._spTree, slide.part, width, height, illustrations=not bookend
        )
        if photos:
            found.append(
                _issue(
                    "integrity.donor_photo_left",
                    slide_ir.index,
                    f"фото шаблона на слайде ({len(photos)}): иллюстрация чужой темы — "
                    "место под картинку, а не содержание",
                    bbox=photos[0],
                )
            )
    return found


def donor_background(pptx_path: str | Path, deck: DeckIR, spec=None) -> list[Issue]:
    """Фон собранного слайда против фона его донора.

    Фон, заданный на слайде шаблона (`p:bg`), а не в мастере, не лежит в
    дереве фигур и при клонировании композиции сам не приезжает: бланк
    «Дорожная карта» собирался белым вместо своей заливки, тёмный финал —
    белым. Сверяется отпечаток разметки фона: заливка, градиент, картинка.
    """
    from pptx import Presentation

    from deckwright.parse.tokens import background_signature, own_background

    patterns = {pattern.id: pattern for pattern in getattr(spec, "patterns", [])}
    found: list[Issue] = []
    presentation = Presentation(str(pptx_path))
    for slide_ir, slide in zip(deck.slides, presentation.slides, strict=False):
        pattern = patterns.get(slide_ir.pattern_id)
        if pattern is None or pattern.background_signature is None:
            continue
        ours = background_signature(own_background(slide._element), slide.part)
        if ours != pattern.background_signature:
            found.append(
                _issue(
                    "template.background_off_donor",
                    slide_ir.index,
                    f"фон не как у донора (слайд {pattern.donor_slide_index} шаблона): "
                    + ("своего фона нет, виден фон мастера" if ours is None else "другая заливка"),
                )
            )
    return found


def run(
    deck: DeckIR,
    plan: DeckPlan,
    pack,
    pptx_path: str | Path | None,
    max_bullets: int,
    max_words: int,
    min_fill: float = MIN_FILL_RATIO,
    spec=None,
) -> list[Issue]:
    found: list[Issue] = []
    for slide in deck.slides:
        found.extend(density(slide, max_bullets, max_words, min_fill))
    found.extend(title_only(deck, spec))
    found.extend(duplicate_slides(deck))
    found.extend(figures(plan, pack))
    found.extend(undeclared_numbers(plan, pack))
    if pptx_path is not None:
        found.extend(fill_ratio(pptx_path, deck, min_fill))
        found.extend(package(pptx_path))
        found.extend(donor_data(pptx_path, deck, spec))
        found.extend(donor_background(pptx_path, deck, spec))
        found.extend(donor_photos(pptx_path, deck, spec))
    return found
