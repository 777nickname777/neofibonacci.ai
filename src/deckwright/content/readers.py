"""Чтение входа: файлы и вставленный текст → текст фрагментами и таблицы.

Слой детерминированный: модель сюда не ходит. Его задача — достать из файла
всё, что в нём написано, с указанием места (страница, слайд, раздел), чтобы
дальше каждое число и каждый факт можно было сверить с источником.

Форматы: `.txt`, `.md`, `.docx`, `.pdf`, `.pptx`, `.json`. JSON в нашей схеме
контент-пакета сюда не попадает — он принимается как есть раньше
(`content.ingest`). JSON чужой структуры разворачивается в строки
«путь: значение»: так его читает модель и так в нём ищутся числа.
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path

# Форматы, которые читаются. Имя — то, что пишется в `SourceDoc.kind`.
KINDS = {
    ".txt": "txt",
    ".md": "md",
    ".markdown": "md",
    ".docx": "docx",
    ".pdf": "pdf",
    ".pptx": "pptx",
    ".json": "json",
}


class UnsupportedInput(ValueError):
    """Файл не того формата или не читается."""


@dataclass
class Fragment:
    """Кусок текста документа и его место: «стр. 2», «слайд 5», «раздел …»."""

    text: str
    locator: str = ""


@dataclass
class Table:
    """Таблица документа: строки ячеек, первая — шапка, если она есть."""

    rows: list[list[str]]
    locator: str = ""


@dataclass
class SourceText:
    """Документ входа, прочитанный целиком."""

    doc_id: str
    name: str
    kind: str
    fragments: list[Fragment] = field(default_factory=list)
    tables: list[Table] = field(default_factory=list)

    @property
    def text(self) -> str:
        """Весь текст документа, включая ячейки таблиц: в нём ищутся числа."""
        parts = [fragment.text for fragment in self.fragments]
        parts += [" | ".join(row) for table in self.tables for row in table.rows]
        return "\n".join(parts)

    @property
    def char_count(self) -> int:
        return len(self.text)


def read_inline(text: str, doc_id: str, name: str = "текст из поля ввода") -> SourceText:
    """Вставленный текст: бриф или готовый материал."""
    fragments, tables = _read_markdown(text)
    return SourceText(doc_id=doc_id, name=name, kind="inline", fragments=fragments, tables=tables)


def read_file(path: str | Path, doc_id: str, name: str | None = None) -> SourceText:
    """Файл входа по расширению."""
    path = Path(path)
    kind = KINDS.get(path.suffix.lower())
    if kind is None:
        supported = ", ".join(sorted(KINDS))
        raise UnsupportedInput(f"{path.name}: формат не поддерживается, нужен один из {supported}")
    name = name or path.name
    try:
        if kind in ("txt", "md"):
            fragments, tables = _read_markdown(_decode(path.read_bytes()))
        elif kind == "docx":
            fragments, tables = _read_docx(path)
        elif kind == "pdf":
            fragments, tables = _read_pdf(path)
        elif kind == "pptx":
            fragments, tables = _read_pptx(path)
        else:
            fragments, tables = _read_json(path.read_bytes())
    except UnsupportedInput:
        raise
    except Exception as exc:  # битый файл любого формата — одна ошибка для UI
        raise UnsupportedInput(f"{name}: не удалось прочитать ({exc})") from exc
    return SourceText(doc_id=doc_id, name=name, kind=kind, fragments=fragments, tables=tables)


def _decode(raw: bytes) -> str:
    """Текстовый файл в незнакомой кодировке: UTF-8, затем cp1251."""
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _read_markdown(text: str) -> tuple[list[Fragment], list[Table]]:
    """Текст и Markdown: абзацы под ближайшим заголовком, таблицы с `|`."""
    fragments: list[Fragment] = []
    tables: list[Table] = []
    section = ""
    paragraph: list[str] = []
    table: list[list[str]] = []

    def flush_paragraph() -> None:
        if paragraph:
            fragments.append(Fragment(" ".join(paragraph), _locator(section, len(fragments))))
            paragraph.clear()

    def flush_table() -> None:
        if table:
            tables.append(Table([row for row in table], _locator(section, len(tables), "таблица")))
            table.clear()

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if line.startswith("|") and line.endswith("|"):
            flush_paragraph()
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            # Строка-разделитель шапки «|---|:--:|» — разметка, не данные.
            if all(set(cell) <= set("-: ") for cell in cells):
                continue
            table.append(cells)
            continue
        flush_table()
        if not line:
            flush_paragraph()
            continue
        if line.startswith("#"):
            flush_paragraph()
            section = line.lstrip("#").strip()
            fragments.append(Fragment(section, section))
            continue
        # Пункт списка — отдельный фрагмент: в брифе это обычно отдельная мысль.
        if line[:2] in ("- ", "* ", "• "):
            flush_paragraph()
            fragments.append(Fragment(line[2:].strip(), _locator(section, len(fragments))))
            continue
        paragraph.append(line)
    flush_paragraph()
    flush_table()
    return fragments, tables


def _locator(section: str, number: int, what: str = "абзац") -> str:
    return f"раздел «{section}»" if section else f"{what} {number + 1}"


def _read_docx(path: Path) -> tuple[list[Fragment], list[Table]]:
    """Абзацы с именем раздела по заголовкам документа и таблицы по порядку."""
    import docx

    document = docx.Document(str(path))
    fragments: list[Fragment] = []
    section = ""
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
        if style.startswith(("heading", "заголовок", "title")):
            section = text
        fragments.append(Fragment(text, f"раздел «{section}»" if section else ""))
    tables = [
        Table(
            [[cell.text.strip() for cell in row.cells] for row in table.rows],
            f"таблица {number + 1}",
        )
        for number, table in enumerate(document.tables)
    ]
    return fragments, [table for table in tables if table.rows]


def _read_pdf(path: Path) -> tuple[list[Fragment], list[Table]]:
    """Текст и таблицы по страницам."""
    import pdfplumber

    fragments: list[Fragment] = []
    tables: list[Table] = []
    with pdfplumber.open(str(path)) as pdf:
        for number, page in enumerate(pdf.pages, start=1):
            locator = f"стр. {number}"
            for table in page.extract_tables() or []:
                rows = [[(cell or "").strip() for cell in row] for row in table]
                rows = [row for row in rows if any(row)]
                if rows:
                    tables.append(Table(rows, locator))
            text = page.extract_text() or ""
            for block in text.split("\n\n"):
                block = " ".join(line.strip() for line in block.splitlines() if line.strip())
                if block:
                    fragments.append(Fragment(block, locator))
    if not fragments and not tables:
        raise UnsupportedInput(f"{path.name}: в PDF нет текстового слоя (скан?)")
    return fragments, tables


def _read_pptx(path: Path) -> tuple[list[Fragment], list[Table]]:
    """Текст фигур и таблицы по слайдам, с заметками докладчика."""
    from pptx import Presentation

    fragments: list[Fragment] = []
    tables: list[Table] = []
    for number, slide in enumerate(Presentation(str(path)).slides, start=1):
        locator = f"слайд {number}"
        for shape in _walk(slide.shapes):
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                tables.append(Table(rows, locator))
            elif shape.has_text_frame and shape.text_frame.text.strip():
                fragments.append(Fragment(" ".join(shape.text_frame.text.split()), locator))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                fragments.append(Fragment(" ".join(notes.split()), f"{locator}, заметки"))
    return fragments, tables


def _walk(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):
            yield from _walk(shape.shapes)


def _read_json(raw: bytes) -> tuple[list[Fragment], list[Table]]:
    """JSON любой структуры → строки «путь: значение».

    Списки одинаковых объектов («kpis»: [{"name": …, "q1": …}]) становятся
    таблицей: так их проще прочитать и модели, и человеку в отчёте, и из них
    сразу получаются ряды для графика.
    """
    data = json.loads(_decode(raw))
    fragments: list[Fragment] = []
    tables: list[Table] = []

    def visit(value, path: str) -> None:
        if isinstance(value, dict):
            for key, item in value.items():
                visit(item, f"{path}.{key}" if path else str(key))
        elif isinstance(value, list):
            if value and all(isinstance(item, dict) for item in value):
                keys = list(dict.fromkeys(key for item in value for key in item))
                scalar = all(
                    not isinstance(item.get(key), (dict, list)) for item in value for key in keys
                )
                if scalar:
                    rows = [keys] + [[_scalar(item.get(key)) for key in keys] for item in value]
                    tables.append(Table(rows, path or "json"))
                    return
            for index, item in enumerate(value):
                visit(item, f"{path}[{index}]")
        elif value is not None and _scalar(value):
            fragments.append(Fragment(f"{path}: {_scalar(value)}", path or "json"))

    visit(data, "")
    return fragments, tables


def _scalar(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "да" if value else "нет"
    return str(value).strip()
