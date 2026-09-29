"""Экспорт: три формата, и каждый несёт то же содержание.

Проверяется не «файл создался», а то, что в нём. PDF из нужного числа страниц,
HTML с настоящим текстом в DOM и без единого внешнего запроса — иначе колода,
открытая без сети, окажется другой.
"""

from __future__ import annotations

import json
import re

import pytest

from deckwright.config import load_config
from deckwright.layout.text_metrics import deck_substitutions
from deckwright.llm.fake import RecordedClient
from deckwright.pipeline import run_variant
from deckwright.render.html import export_html
from deckwright.schemas import ContentPack

CONFIG = "configs/config.yaml"


@pytest.fixture(scope="module")
def pack(content_pack_path):
    return ContentPack.model_validate(json.loads(content_pack_path.read_text("utf-8")))


@pytest.fixture(scope="module")
def exported(template_paths, pack, recorded_dir, tmp_path_factory):
    cfg = load_config(CONFIG)
    root = tmp_path_factory.mktemp("export")
    return [
        (
            path.stem,
            run_variant(
                template_path=path,
                pack=pack,
                cfg=cfg,
                client=RecordedClient(recorded_dir),
                variant="balanced",
                output_dir=root / path.stem,
            ),
        )
        for path in template_paths
    ]


def test_all_three_formats_are_produced(exported):
    """ТЗ просит .pptx, .pdf и .html — и все три обязаны существовать."""
    for name, result in exported:
        assert result.pptx.exists(), f"{name}: нет .pptx"
        assert result.pdf.exists(), f"{name}: нет .pdf"
        assert result.html.exists(), f"{name}: нет .html"
        assert result.pages, f"{name}: нет картинок слайдов"


def test_pdf_has_a_page_per_slide(exported):
    """Расхождение значит, что конвертация потеряла или удвоила слайд.

    Считается по картинкам: `pdftoppm` делает ровно одну на страницу, и это
    честнее, чем верить заголовку PDF.
    """
    for name, result in exported:
        assert len(result.pages) == len(result.deck.slides), (
            f"{name}: страниц {len(result.pages)}, слайдов {len(result.deck.slides)}"
        )


def test_html_carries_real_text_not_a_picture_of_it(exported):
    """Соблазн был конвертировать PDF в SVG — замер показал, чего это стоит.

    На одном слайде `vk_workspace` выходило 855 КБ, ноль элементов `<text>` и
    338 ссылок на глифы: весь текст переведён в кривые. Ни выделить, ни найти
    поиском. Здесь текст обязан быть текстом.
    """
    for name, result in exported:
        page = result.html.read_text("utf-8")
        titles = [slide.takeaway_title for slide in result.plan.slides]
        landed = [title for title in titles if title and title in page]
        assert landed, f"{name}: ни один заголовок плана не доехал до HTML"
        assert page.count("<section") == len(result.deck.slides), (
            f"{name}: секций в HTML не столько, сколько слайдов"
        )


def test_html_is_self_contained(exported):
    """Ни одного внешнего запроса: без сети страница обязана выглядеть так же.

    Ссылка на шрифт или картинку в интернете означает, что колода у проверяющего
    и колода у нас — разные документы.
    """
    for name, result in exported:
        page = result.html.read_text("utf-8")
        external = re.findall(r'(?:src|href)\s*=\s*"(?!data:|#)([^"]+)"', page)
        assert not external, f"{name}: HTML тянет внешние ресурсы: {external[:3]}"


def test_html_keeps_the_template_font_names(exported):
    """Имя гарнитуры в HTML то же, что в шаблоне и в `.pptx`.

    Подменять его нельзя нигде: человек, открывший страницу, должен видеть
    шрифт шаблона, а аудит — не находить чужой.
    """
    for name, result in exported:
        page = result.html.read_text("utf-8")
        families = {token.family for token in result.spec.fonts}
        if not families:
            continue
        assert any(f"'{family}'" in page for family in families), (
            f"{name}: в HTML нет ни одной гарнитуры шаблона ({sorted(families)})"
        )


def test_html_embeds_each_picture_once(exported):
    """Декор шаблона лежит на мастере и повторяется на каждом слайде.

    Вкладывать его по разу на слайд — тот же файл двенадцать раз: страница
    `vk_workspace` выходила 31 МБ вместо восьми.
    """
    for name, result in exported:
        page = result.html.read_text("utf-8")
        blobs = re.findall(r"base64,([A-Za-z0-9+/=]{200,})", page)
        assert len(blobs) == len(set(blobs)), f"{name}: одни и те же байты в HTML дважды"


def test_html_without_a_deck_file_still_renders_text(deck_and_spec, tmp_path):
    """Страница обязана собираться и без `.pptx`: картинок не будет, текст будет.

    Это тот случай, когда экспорт зовут по одному `SlideIR` — например, чтобы
    посмотреть вёрстку до сборки колоды.
    """
    deck, spec = deck_and_spec
    path = export_html(deck, tmp_path / "bare.html", spec=spec, pptx_path=None)
    page = path.read_text("utf-8")
    assert page.count("<section") == len(deck.slides)
    assert "<img" not in page and "background-image" not in page


@pytest.fixture(scope="module")
def deck_and_spec(template_paths, pack, recorded_dir):
    from deckwright.layout.matcher import build_deck_ir
    from deckwright.parse.opener import parse_template
    from deckwright.plan.planner import build_plan

    cfg = load_config(CONFIG)
    spec = parse_template(template_paths[0])
    plan, _, _ = build_plan(pack, RecordedClient(recorded_dir), cfg.deck.min_slides)
    deck, _ = build_deck_ir(spec, plan, cfg.variant("balanced"), pack=pack)
    return deck, spec


# ── Частичная растеризация ───────────────────────────────────────────────────
#
# Итерация цикла исправления трогает два-три слайда, а растеризация всей
# колоды стоит 18 с из 20. Перерисовывать все страницы ради трёх незачем — но
# переиспользование прошлых картинок обязано быть честным: страница с номером
# 4 после правки может оказаться другим слайдом.


def test_only_requested_pages_are_rerendered(exported, tmp_path):
    """Перерисовывается только названная страница, список остаётся полным."""
    from deckwright.render.png import pdf_to_png

    result = exported[0][1]
    directory = tmp_path / "png"
    first = pdf_to_png(result.pdf, directory, dpi=48)
    assert len(first) == len(result.deck.slides)

    marks = {page: page.stat().st_mtime_ns for page in first}
    again = pdf_to_png(result.pdf, directory, dpi=48, only_pages={2})

    assert len(again) == len(first), "список страниц обязан остаться полным"
    changed = [page for page in again if page.stat().st_mtime_ns != marks[page]]
    assert [page.name for page in changed] == [first[1].name]


def test_empty_selection_rerenders_nothing(exported, tmp_path):
    """Правка могла не тронуть ни одной страницы — это законный случай."""
    from deckwright.render.png import pdf_to_png

    result = exported[0][1]
    directory = tmp_path / "png"
    first = pdf_to_png(result.pdf, directory, dpi=48)
    marks = {page: page.stat().st_mtime_ns for page in first}

    again = pdf_to_png(result.pdf, directory, dpi=48, only_pages=set())
    assert len(again) == len(first)
    assert all(page.stat().st_mtime_ns == marks[page] for page in again)


def test_shorter_deck_leaves_no_stale_pages(exported, tmp_path):
    """Колода стала короче — хвост прошлой сборки в отчёт попасть не должен.

    Иначе аудит получил бы картинку слайда, которого в `.pdf` уже нет, а
    интерфейс показал бы его пользователю.
    """
    from deckwright.render.png import pdf_to_png

    result = exported[0][1]
    directory = tmp_path / "png"
    pages = pdf_to_png(result.pdf, directory, dpi=48)
    stale = directory / f"{result.pdf.stem}-{len(pages) + 1:02d}.png"
    stale.write_bytes(pages[0].read_bytes())

    # Число картинок разошлось с числом страниц: частичная растеризация
    # запрещена, идёт полная, и лишний файл убирается.
    again = pdf_to_png(result.pdf, directory, dpi=48, only_pages={1})
    assert not stale.exists(), "осталась страница от прошлой, более длинной сборки"
    assert len(again) == len(pages)


def test_pdf_is_drawn_in_the_templates_own_font(exported):
    """Картинка рисуется тем шрифтом, которым фиттер мерил текст.

    Без этого `vk_tech` рисовался DejaVu Sans вместо Play: шире, и текст,
    честно уложенный по метрикам Play, на картинке рвал слова посередине.

    Сверяются гарнитуры, которыми колода **набрана**: фиттер меряет их же
    или метрически совместимым клоном, а «любой встроенный в шаблон шрифт»
    — не то же самое. В `vk_education` встроен Play, но текст набран Arial,
    и требовать Play в его PDF не за что.
    """
    import shutil
    import subprocess

    if shutil.which("pdffonts") is None:
        pytest.skip("pdffonts не установлен")
    checked = 0
    for name, result in exported:
        embedded = {
            token.family.replace(" ", "")
            for token in result.spec.fonts
            if token.embedded and token.file_path
        }
        if not embedded:
            # Шрифта шаблона в файле нет: в системе его может не быть тоже, и
            # подстановка тогда — не дефект, а единственный выход.
            continue
        cfg = load_config(CONFIG)
        # Гарнитура, которой шрифта нет ни в шаблоне, ни в системе, честно
        # заменяется — и замена объявлена конвертеру. Тогда в PDF обязана
        # стоять она, а не что-нибудь третье.
        swapped = {
            item["requested"].replace(" ", ""): item["used"].replace(" ", "")
            for item in deck_substitutions(
                result.deck, result.spec, cfg.fonts.substitution_slack
            )
        }
        families = {
            paragraph.style.font_family.replace(" ", "")
            for slide in result.deck.slides
            for element in slide.all_elements()
            if element.text is not None
            for paragraph in element.text.paragraphs
            if paragraph.style.font_family
        }
        listing = subprocess.run(
            ["pdffonts", str(result.pdf)], capture_output=True, text=True, check=True
        ).stdout.replace(" ", "")
        missing = sorted(
            family
            for family in families
            if swapped.get(family, family) not in listing
        )
        assert not missing, (
            f"{name}: набрано {missing}, замены {swapped}, в PDF {listing[:300]}"
        )
        checked += 1
    if not checked:
        pytest.skip("ни в одном шаблоне нет встроенных шрифтов")


# ── Номера элементов повторителя ─────────────────────────────────────────────


def _ordinal_case():
    """Геометрия кружков с номерами из живого шаблона `finansy`.

    Слот номера снят с первого кружка; шаг повторителя снят с карточек под
    ними. Дизайнер поставил кружки на глаз, и второй с третьим разошлись со
    своими расчётными местами на 0.026″ и 0.04″ — больше допуска точного
    совпадения рамки (0.01″).
    """
    from deckwright.schemas import Box

    slot = Box(x=1_455_966, y=1_714_452, w=733_424, h=707_886)
    step = (0, 3_373_437, 6_797_676)
    real = [
        Box(x=1_455_966, y=1_714_452, w=733_424, h=707_886),
        Box(x=4_829_403, y=1_690_688, w=733_424, h=707_886),
        Box(x=8_280_176, y=1_678_047, w=733_424, h=707_886),
    ]
    return slot, step, real


def test_ordinal_shape_is_found_though_it_stands_by_eye():
    """Номер элемента находится, даже если кружок стоит не по шагу.

    Без этого на `finansy` (run 34) номера 2 и 3 не находились, дальше их
    стирало как числа донора, и на слайде оставался номер «1» при двух
    пустых кружках.
    """
    from deckwright.render import pptx_writer as writer
    from deckwright.schemas import Box

    slot, step, real = _ordinal_case()
    candidates = [(box, f"кружок {index + 1}") for index, box in enumerate(real)]
    found = []
    for index, dx in enumerate(step):
        want = slot.model_copy(update={"x": slot.x + dx})
        band = Box(x=real[index].x - 400_000, y=0, w=3_000_000, h=6_858_000)
        found.append(writer._take_ordinal_shape(candidates, want, band))
    assert found == ["кружок 1", "кружок 2", "кружок 3"]
    assert not candidates, "найденная фигура обязана выниматься из списка"


def test_ordinal_search_keeps_its_limits():
    """Поиск на глаз не хватает что попало: размер, полоса и предел сдвига."""
    from deckwright.render import pptx_writer as writer
    from deckwright.schemas import Box

    slot, _, real = _ordinal_case()
    band = Box(x=real[1].x - 400_000, y=0, w=3_000_000, h=6_858_000)
    want = slot.model_copy(update={"x": real[1].x})

    # Подпись карточки — другого размера, номером не станет.
    caption = Box(x=real[1].x, y=real[1].y, w=2_808_000, h=707_886)
    assert writer._take_ordinal_shape([(caption, "подпись")], want, band) is None

    # Номер соседнего элемента лежит в чужой полосе.
    assert writer._take_ordinal_shape([(real[2], "чужой")], want, band) is None

    # Сдвиг больше 0.1″ — это уже не «на глаз», а другая фигура.
    far = real[1].model_copy(update={"y": real[1].y + 200_000})
    assert writer._take_ordinal_shape([(far, "далеко")], want, band) is None

    # А в пределах допуска — находится.
    near = real[1].model_copy(update={"y": real[1].y + 40_000})
    assert writer._take_ordinal_shape([(near, "рядом")], want, band) == "рядом"


def test_donor_frame_is_found_though_it_stands_by_eye():
    """Рамка донора, сдвинутая на глаз, — та же рамка, а не новая надпись.

    Шаг повторителя снят с подложек карточек, а надписи внутри дизайнер
    двигал отдельно: на `finansy` подписи второй и третьей карточки стоят
    на 0.03″ от расчётного места. Рамка не находилась, текст уезжал в новую
    надпись — без маркера списка и прочего оформления донора, — а пустая
    рамка донора оставалась на слайде.
    """
    from deckwright.render import pptx_writer as writer
    from deckwright.schemas import Box

    want = Box(x=3_776_663, y=2_486_025, w=2_808_000, h=1_000_000)
    donor = want.model_copy(update={"x": want.x - 27_432})  # 0.03″ левее

    candidates = [(donor, "рамка донора")]
    assert writer._take_matching_shape(candidates, want) == "рамка донора"
    assert not candidates

    # Другой размер — другое место, а не сдвинутое.
    other = want.model_copy(update={"x": want.x - 27_432, "w": want.w // 2})
    assert writer._take_matching_shape([(other, "чужая")], want) is None

    # Сдвиг больше 0.05″ — тоже другое место.
    far = want.model_copy(update={"x": want.x - 91_440})
    assert writer._take_matching_shape([(far, "далеко")], want) is None


def _element_for(frame, bullets: list[bool]):
    """Элемент представления под уже набранную рамку."""
    from deckwright.schemas import (
        Box,
        Color,
        Element,
        ElementKind,
        Paragraph,
        Provenance,
        SlotRole,
        SourceKind,
        TextContent,
        TextStyle,
    )

    style = TextStyle(font_family="Arial", size_pt=18, color=Color(rgb="111111"))
    return Element(
        id="e",
        kind=ElementKind.TEXT,
        role=SlotRole.BODY,
        box=Box(x=0, y=0, w=1_000_000, h=500_000),
        provenance=Provenance(kind=SourceKind.DERIVED, ref="тест"),
        text=TextContent(
            paragraphs=[
                Paragraph(text=p.text or "x", style=style, bullet=flag)
                for p, flag in zip(frame.paragraphs, bullets, strict=True)
            ]
        ),
    )


def _frame_with_bullet(bullet: str):
    """Надпись, первому абзацу которой задан маркер, и два абзаца после него."""
    from pptx import Presentation
    from pptx.util import Emu

    from deckwright.render.pptx_writer import _A

    slide = Presentation().slides.add_slide(Presentation().slide_layouts[6])
    frame = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(2_000_000), Emu(1_000_000)).text_frame
    first = frame.paragraphs[0]
    first.add_run().text = "Июль"
    ppr = first._p.get_or_add_pPr()
    ppr.append(ppr.makeelement(f"{_A}{bullet}", {}))
    for text in ("Август", "Сентябрь"):
        frame.add_paragraph().add_run().text = text
    return frame


def _bullets_of(frame):
    from lxml import etree

    from deckwright.render.pptx_writer import _A

    marks = []
    for paragraph in frame.paragraphs:
        ppr = paragraph._p.find(f"{_A}pPr")
        names = [etree.QName(node).localname for node in ppr] if ppr is not None else []
        marks.append([name for name in names if name.startswith("bu")])
    return marks


def test_added_paragraphs_take_the_bullet_of_the_first():
    """Список выглядит списком: маркер у всех пунктов одинаковый.

    Донор задаёт оформление своему единственному абзацу — на `finansy` это
    «маркера нет». Абзацы, которые дописываем мы, брали маркер с layout'а,
    и на слайде выходило «Июль» без точки, «Август» и «Сентябрь» с точками.
    """
    from deckwright.render.pptx_writer import _align_bullets

    frame = _frame_with_bullet("buNone")
    assert _bullets_of(frame) == [["buNone"], [], []], "проба собрана неверно"
    _align_bullets(frame, _element_for(frame, [False, False, False]))
    assert _bullets_of(frame) == [["buNone"], ["buNone"], ["buNone"]]


def test_a_bullet_of_the_template_reaches_every_paragraph():
    """Обратный случай: донор с маркером отдаёт его всем абзацам."""
    from deckwright.render.pptx_writer import _align_bullets

    frame = _frame_with_bullet("buChar")
    _align_bullets(frame, _element_for(frame, [True, True, True]))
    assert _bullets_of(frame) == [["buChar"], ["buChar"], ["buChar"]]


def test_a_frame_without_its_own_bullet_is_left_alone():
    """Донор ничего не задавал — не выдумываем за него."""
    from pptx import Presentation
    from pptx.util import Emu

    from deckwright.render.pptx_writer import _align_bullets

    slide = Presentation().slides.add_slide(Presentation().slide_layouts[6])
    frame = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(2_000_000), Emu(1_000_000)).text_frame
    frame.paragraphs[0].add_run().text = "Июль"
    frame.add_paragraph().add_run().text = "Август"
    _align_bullets(frame, _element_for(frame, [True, True]))
    assert _bullets_of(frame) == [[], []]


def test_a_single_line_gets_no_bullet_from_the_donor():
    """Не список — значит без маркера, даже если карточка донора с ним.

    Маркер приносит свой отступ, и рамке перестаёт хватать ширины: на
    `finansy` из-за него LibreOffice ломал «Автоматическая» посреди слова,
    хотя фиттер мерил строку без отступа.
    """
    from deckwright.render.pptx_writer import _align_bullets

    frame = _frame_with_bullet("buChar")
    _align_bullets(frame, _element_for(frame, [False, False, False]))
    assert _bullets_of(frame) == [["buNone"], ["buNone"], ["buNone"]]


# ── Рендер ничего не выбирает сам ────────────────────────────────────────────


def test_every_branch_of_the_writer_puts_the_resolved_style_in_the_file(exported):
    """Гарнитура, кегль и цвет в `.pptx` — те же, что в `SlideIR`.

    Веток записи текста три: фигура донора, плейсхолдер макета и новая
    надпись. Плейсхолдер раньше не получал типографики вовсе — «оставим
    шаблонную», — и гарнитура молча доставалась теме: заголовок мерился
    `Calibri`, а рисовался `Calibri Light`.
    """
    from pptx import Presentation

    from deckwright.schemas import ElementKind

    def walk(shapes):
        for shape in shapes:
            yield shape
            if hasattr(shape, "shapes"):
                yield from walk(shape.shapes)

    checked = 0
    for name, result in exported:
        deck = Presentation(str(result.pptx))
        for slide, ir in zip(deck.slides, result.deck.slides, strict=True):
            by_text = {}
            for shape in walk(slide.shapes):
                if shape.has_text_frame and shape.text_frame.text.strip():
                    by_text.setdefault(shape.text_frame.text.strip(), shape)
            for element in ir.all_elements():
                if element.kind is not ElementKind.TEXT or element.text is None:
                    continue
                key = "\n".join(p.text for p in element.text.paragraphs).strip()
                shape = by_text.get(key)
                if shape is None:
                    continue
                runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
                if not runs:
                    continue
                want = element.text.paragraphs[0].style
                checked += 1
                assert runs[0].font.name == want.font_family, (
                    f"{name}, слайд {ir.index}, {element.id}: в файле "
                    f"{runs[0].font.name!r}, в IR {want.font_family!r}"
                )
                assert runs[0].font.size is not None, (
                    f"{name}, слайд {ir.index}, {element.id}: кегль не записан"
                )
                # Допуск — сотая пункта: python-pptx хранит кегль в
                # сотых долях, и 8.12 pt возвращается как 8.11.
                assert runs[0].font.size.pt == pytest.approx(want.size_pt, abs=0.02), (
                    f"{name}, слайд {ir.index}, {element.id}: кегль "
                    f"{runs[0].font.size.pt} против {want.size_pt}"
                )
    assert checked, "ни одного текстового элемента не сверили"


def test_the_deck_writes_the_fonts_of_its_template(exported):
    """Ни одной гарнитуры со стороны: всё, чем написана колода, — из шаблона."""
    from deckwright.layout.text_metrics import normalized_family

    for name, result in exported:
        known = {normalized_family(token.family) for token in result.spec.fonts}
        used = {
            normalized_family(p.style.font_family)
            for slide in result.deck.slides
            for element in slide.all_elements()
            if element.text is not None
            for p in element.text.paragraphs
        }
        assert used <= known, f"{name}: чужие гарнитуры {sorted(used - known)}"
