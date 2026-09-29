"""Оформление шаблона участвует в размещении.

Проверки поведения, а не устройства: шаблон собирается фикстурой с
композицией, которой нет в датасете, и ни один тест не смотрит на имена
`vk_tech`/`zelenie` — решение должно быть общим.

До этих изменений вёрстка знала только про фигуры самого слайда. Оформление
обычно живёт в макете: у одного из шаблонов датасета это 160 картинок мимо
всех ограничений, и текст ложился поверх них.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).parent / "fixtures"))

from deckwright.config import load_config  # noqa: E402
from deckwright.layout.matcher import build_deck_ir, content_regions  # noqa: E402
from deckwright.parse import surfaces  # noqa: E402
from deckwright.parse.opener import parse_template  # noqa: E402
from deckwright.schemas import ContentPack, DeckPlan  # noqa: E402

IN = 914400


@pytest.fixture(scope="module")
def decor_template(tmp_path_factory) -> Path:
    from make_decor_template import build_decor_template

    target = tmp_path_factory.mktemp("decor") / "template.pptx"
    return build_decor_template(target)


@pytest.fixture(scope="module")
def decor_spec(decor_template):
    return parse_template(decor_template, cache_dir=None, font_dir=None)


def _pattern_with_photo(spec):
    """Композиция, под которой лежит полноразмерная фотография макета."""
    for pattern in spec.patterns:
        for box in pattern.layout_obstacles:
            if box.w >= spec.slide_width_emu * 0.9 and box.h >= spec.slide_height_emu * 0.9:
                return pattern
    raise AssertionError("в фикстуре нет композиции с полноразмерным фото")


def _pattern_with_side_image(spec):
    """Композиция «картинка сбоку»: препятствие занимает часть ширины."""
    for pattern in spec.patterns:
        for box in pattern.layout_obstacles:
            if 0.2 < box.w / spec.slide_width_emu < 0.6 and box.h > spec.slide_height_emu * 0.5:
                return pattern
    raise AssertionError("в фикстуре нет композиции с боковой картинкой")


# ── оформление доходит до ограничений ───────────────────────────────────


def test_layout_artwork_reaches_the_constraints(decor_spec):
    """Фигуры макета видны вёрстке, а не только самому макету."""
    assert any(layout.decor_boxes for layout in decor_spec.layouts), (
        "ни один макет не отдал оформление"
    )
    assert any(pattern.layout_obstacles for pattern in decor_spec.patterns), (
        "оформление макета не дошло до композиций"
    )


def test_logo_in_the_layout_is_protected(decor_spec):
    """Значки логотипа из макета — препятствие, а не свободное место."""
    pattern = _pattern_with_photo(decor_spec)
    marks = [
        box for box in pattern.layout_obstacles
        if box.w <= int(0.5 * IN) and box.h <= int(0.5 * IN)
    ]
    assert len(marks) >= 2, f"значки логотипа не защищены: {pattern.layout_obstacles}"
    for region in content_regions(decor_spec, pattern):
        for mark in marks:
            assert region.intersection(mark) is None, "свободная область накрыла логотип"


# ── фон, фотография и подложка ──────────────────────────────────────────


def test_full_bleed_photo_is_not_a_safe_background(decor_spec):
    """Фотография во весь слайд остаётся фотографией.

    Площадь сама по себе роли не определяет: по ровной заливке текст писать
    можно, по фотографии — нет, сколько бы места она ни занимала.
    """
    pattern = _pattern_with_photo(decor_spec)
    full = [
        box for box in pattern.layout_obstacles
        if box.w >= decor_spec.slide_width_emu * 0.9
    ]
    assert full, "полноразмерное фото не признано препятствием"


def test_designed_text_panel_gives_the_space_back(decor_spec):
    """Подложка поверх фотографии — это приглашение писать именно там."""
    pattern = _pattern_with_photo(decor_spec)
    regions = content_regions(decor_spec, pattern)
    assert regions, "слайд с фото не принял бы ни одной надписи"
    panel = max(pattern.layout_backdrops, key=lambda b: b.w * b.h)
    assert any(
        region.intersection(panel) is not None for region in regions
    ), "свободное место не совпало с подложкой"
    # И при этом не растеклось на всю фотографию.
    biggest = max(regions, key=lambda b: b.w * b.h)
    assert biggest.w * biggest.h <= panel.w * panel.h * 1.05


def test_background_does_not_block_the_whole_slide(decor_spec):
    """Ни одна композиция не остаётся вовсе без места."""
    for pattern in decor_spec.patterns:
        assert content_regions(decor_spec, pattern), f"{pattern.id}: места нет совсем"


# ── несколько свободных областей ────────────────────────────────────────


def test_side_image_leaves_a_column_beside_it(decor_spec):
    """Картинка справа оставляет колонку слева, а не «одну область на всё»."""
    pattern = _pattern_with_side_image(decor_spec)
    picture = next(
        box for box in pattern.layout_obstacles
        if 0.2 < box.w / decor_spec.slide_width_emu < 0.6
    )
    regions = content_regions(decor_spec, pattern)
    assert len(regions) >= 2, f"сложный макет сведён к одной рамке: {regions}"
    for region in regions:
        assert region.intersection(picture) is None, "область наложилась на картинку"
    column = max(regions, key=lambda b: b.w * b.h)
    assert column.right <= picture.x + IN // 100, "колонка заехала под картинку"


# ── наследование и геометрия ────────────────────────────────────────────


def test_inherited_artwork_is_counted_once(decor_template):
    """Фигура образца, повторённая макетом, не удваивается."""
    from pptx import Presentation

    prs = Presentation(str(decor_template))
    master = prs.slide_masters[0]
    for layout in master.slide_layouts:
        found = surfaces.collect(
            layout._element, layout.part, master._element, master.part,
            prs.slide_width, prs.slide_height,
        )
        boxes = [(s.box.x, s.box.y, s.box.w, s.box.h) for s in found]
        assert len(boxes) == len(set(boxes)), f"{layout.name}: оформление удвоено"


def test_group_is_counted_by_its_own_extent(decor_template):
    """Группа занимает место за себя, а не за себя и каждого ребёнка."""
    from pptx import Presentation

    prs = Presentation(str(decor_template))
    master = prs.slide_masters[0]
    for layout in master.slide_layouts:
        found = surfaces.collect(
            layout._element, layout.part, master._element, master.part,
            prs.slide_width, prs.slide_height,
        )
        total = sum(s.box.w * s.box.h for s in found)
        # Сумма площадей может превышать слайд при честных наложениях, но не
        # в разы: кратное превышение означает, что дети посчитаны вместе с
        # родителем.
        assert total <= 4 * prs.slide_width * prs.slide_height, (
            f"{layout.name}: площадь посчитана многократно"
        )


def test_rotated_element_reports_the_space_it_occupies(decor_spec):
    """Повёрнутая фигура учитывается объемлющим прямоугольником."""
    pattern = _pattern_with_side_image(decor_spec)
    boxes = list(pattern.layout_obstacles) + list(pattern.layout_backdrops)
    # Штрих фикстуры нарисован 3.00x0.06in и повёрнут на 8°: его габарит
    # обязан стать выше исходных 0.06in.
    strokes = [b for b in boxes if 2.5 * IN <= b.w <= 3.5 * IN]
    assert strokes, f"повёрнутый штрих не найден: {[(b.w / IN, b.h / IN) for b in boxes]}"
    assert max(b.h for b in strokes) > int(0.06 * IN) * 1.5, "поворот не учтён"


# ── кэш ─────────────────────────────────────────────────────────────────


def test_cached_and_fresh_parses_agree(decor_template, tmp_path):
    """Из кэша приходят те же ограничения, что при свежем разборе."""
    cache = tmp_path / "cache"
    fresh = parse_template(decor_template, cache_dir=None, font_dir=None)
    first = parse_template(decor_template, cache_dir=cache, font_dir=None)
    second = parse_template(decor_template, cache_dir=cache, font_dir=None)
    counts = [
        sum(len(p.layout_obstacles) for p in spec.patterns)
        for spec in (fresh, first, second)
    ]
    assert len(set(counts)) == 1, f"разбор и кэш разошлись: {counts}"
    assert first.model_dump_json() == second.model_dump_json()


def test_stale_cache_without_constraints_is_not_reused(decor_template, tmp_path):
    """Кэш, записанный до появления ограничений, не подставляется молча."""
    cache = tmp_path / "cache"
    spec = parse_template(decor_template, cache_dir=cache, font_dir=None)
    written = list(cache.glob("*.json"))
    assert written, "кэш не записался"
    # Подменяем содержимое кэша устаревшим: без ограничений макета.
    stale = json.loads(written[0].read_text("utf-8"))
    for pattern in stale.get("patterns", []):
        pattern["layout_obstacles"] = []
    # Имя файла несёт отпечаток разбора и схемы: старое имя другое.
    old_name = written[0].with_name(written[0].name.replace(
        written[0].stem.split("-")[-1], "0" * 12))
    old_name.write_text(json.dumps(stale), encoding="utf-8")
    again = parse_template(decor_template, cache_dir=cache, font_dir=None)
    assert sum(len(p.layout_obstacles) for p in again.patterns) == sum(
        len(p.layout_obstacles) for p in spec.patterns
    ), "подставился кэш без ограничений"


# ── размещение на этом шаблоне ──────────────────────────────────────────


def test_text_does_not_land_on_layout_imagery(decor_template):
    """Ни один текст колоды не ложится на картинку макета."""
    cfg = load_config(ROOT / "configs" / "config.yaml")
    spec = parse_template(decor_template, cache_dir=None, font_dir=cfg.fonts.extract_dir)
    plan = DeckPlan.model_validate(
        json.loads((ROOT / "tests/fixtures/recorded/plan_deck.json").read_text("utf-8"))
    )
    pack = ContentPack.model_validate(
        json.loads((ROOT / "tests/fixtures/content_pack.json").read_text("utf-8"))
    )
    deck, _ = build_deck_ir(spec, plan, cfg.variants[1], pack=pack)

    by_id = {p.id: p for p in spec.patterns}
    offenders = []
    for slide in deck.slides:
        pattern = by_id.get(slide.pattern_id)
        if pattern is None:
            continue
        safe = list(pattern.layout_backdrops)
        for element in slide.all_elements():
            if element.text is None:
                continue
            for picture in pattern.layout_obstacles:
                overlap = element.box.intersection(picture)
                if overlap is None or overlap.area < element.box.area * 0.5:
                    continue
                # Текст на подложке поверх картинки — это замысел шаблона.
                if any(
                    (element.box.intersection(panel) or picture).area
                    >= element.box.area * 0.5
                    for panel in safe
                ):
                    continue
                offenders.append((slide.index, element.id))
    assert not offenders, f"текст лёг на оформление макета: {offenders}"


def test_content_is_not_dropped_when_a_bookend_has_no_room(decor_template):
    """Блок, которому нет места на титуле, не исчезает из колоды."""
    cfg = load_config(ROOT / "configs" / "config.yaml")
    spec = parse_template(decor_template, cache_dir=None, font_dir=cfg.fonts.extract_dir)
    plan = DeckPlan.model_validate(
        json.loads((ROOT / "tests/fixtures/recorded/plan_deck.json").read_text("utf-8"))
    )
    pack = ContentPack.model_validate(
        json.loads((ROOT / "tests/fixtures/content_pack.json").read_text("utf-8"))
    )
    deck, _issues = build_deck_ir(spec, plan, cfg.variants[1], pack=pack)

    written = " ".join(
        paragraph.text
        for slide in deck.slides
        for element in slide.all_elements()
        if element.text is not None
        for paragraph in element.text.paragraphs
    )
    missing = []
    for slide_plan in plan.slides:
        for block in slide_plan.blocks:
            for item in block.items:
                if item and item[:24] not in written:
                    missing.append(item[:40])
    assert not missing, f"содержание потеряно: {missing[:4]}"


def test_insufficient_room_is_reported_not_silent(decor_spec):
    """Нехватка места объявляется находкой, а не молчаливой обрезкой."""
    from deckwright.audit.registry import BY_ID

    assert "layout.text_without_place" in BY_ID
    assert BY_ID["layout.text_without_place"].fix.value in ("assisted", "automatic")


def test_fixed_slide_count_is_respected(decor_template):
    """Заданное пользователем число слайдов вёрстка не увеличивает.

    Без ограничения блок, которому нет места на титуле, уезжает на новый
    слайд — это правильно, когда число слайдов выбирает планировщик. Но
    если число задано, добавить некуда: нехватка места остаётся названной
    находкой, а колода остаётся той длины, которую просил человек.
    """
    cfg = load_config(ROOT / "configs" / "config.yaml")
    spec = parse_template(decor_template, cache_dir=None, font_dir=cfg.fonts.extract_dir)
    plan = DeckPlan.model_validate(
        json.loads((ROOT / "tests/fixtures/recorded/plan_deck.json").read_text("utf-8"))
    )
    pack = ContentPack.model_validate(
        json.loads((ROOT / "tests/fixtures/content_pack.json").read_text("utf-8"))
    )

    free, _ = build_deck_ir(spec, plan, cfg.variants[1], pack=pack, allow_extra_slides=True)
    fixed, issues = build_deck_ir(
        spec, plan, cfg.variants[1], pack=pack, allow_extra_slides=False
    )

    assert len(fixed.slides) == len(plan.slides), (
        f"вёрстка изменила число слайдов: {len(fixed.slides)} против {len(plan.slides)}"
    )
    assert len(free.slides) >= len(fixed.slides)
    # Невозможность размещения названа, а не проглочена.
    assert any(
        issue.check_id in ("layout.text_without_place", "layout.text_overflow")
        for issue in issues
    ), "при фиксированном числе слайдов нехватка места не объявлена"


def test_extra_slides_stop_at_the_ceiling(decor_template):
    """Лишний слайд не переводит колоду за верхнюю границу ТЗ.

    План из 13 слайдов и три переноса давали на `finansy` колоду в 16
    слайдов при требовании «10–15»: потолок `deck.max_slides` вёрстка не
    перепроверяла вовсе — решала только «задано ли точное число».

    Потолок уступает одному — слайду, который нельзя прочитать. График,
    которому досталась полоса в 1 EMU, колоду не спасает тем, что она
    короче, поэтому «не длиннее потолка» проверяется вместе с «ни одного
    сплющенного графика», а не вместо него.
    """
    cfg = load_config(ROOT / "configs" / "config.yaml")
    spec = parse_template(decor_template, cache_dir=None, font_dir=cfg.fonts.extract_dir)
    plan = DeckPlan.model_validate(
        json.loads((ROOT / "tests/fixtures/recorded/plan_deck.json").read_text("utf-8"))
    )
    pack = ContentPack.model_validate(
        json.loads((ROOT / "tests/fixtures/content_pack.json").read_text("utf-8"))
    )
    planned = len(plan.slides)

    def cramped(deck):
        return [
            (slide.index, element.box.h)
            for slide in deck.slides
            for element in slide.all_elements()
            if element.chart is not None
            and element.box.h < deck.slide_height_emu // 4
        ]

    free, _ = build_deck_ir(spec, plan, cfg.variants[1], pack=pack)
    assert len(free.slides) > planned, (
        "проба бессмысленна: без потолка вёрстка слайдов не добавляет"
    )

    capped, issues = build_deck_ir(
        spec, plan, cfg.variants[1], pack=pack, max_slides=planned
    )
    assert len(capped.slides) < len(free.slides), "потолок не укоротил колоду"
    assert not cramped(capped), "потолок оставил график, который не прочитать"
    assert any(
        issue.check_id in ("layout.text_without_place", "layout.text_overflow")
        for issue in issues
    ), "при потолке нехватка места не объявлена"

    # Потолок выше, чем колода хочет: он ничего не меняет.
    loose, _ = build_deck_ir(
        spec, plan, cfg.variants[1], pack=pack, max_slides=planned + 50
    )
    assert len(loose.slides) == len(free.slides)


def test_the_ceiling_is_kept_exactly_when_the_deck_stays_readable():
    """Когда сжатый вариант читается, потолок соблюдается ровно.

    Шаблон-бланк `dorozhnaya_karta`: без потолка план из 10 слайдов растёт
    до 13, и каждый лишний слайд — перенос текста, а не спасение графика.
    """
    template = ROOT / "data/holdout/dorozhnaya_karta.pptx"
    if not template.exists():
        pytest.skip("нет шаблона dorozhnaya_karta.pptx")
    cfg = load_config(ROOT / "configs" / "config.yaml")
    spec = parse_template(template, cache_dir=None, font_dir=cfg.fonts.extract_dir)
    plan = DeckPlan.model_validate(
        json.loads((ROOT / "tests/fixtures/recorded/plan_deck.json").read_text("utf-8"))
    )
    pack = ContentPack.model_validate(
        json.loads((ROOT / "tests/fixtures/content_pack.json").read_text("utf-8"))
    )
    planned = len(plan.slides)

    free, _ = build_deck_ir(spec, plan, cfg.variants[1], pack=pack)
    assert len(free.slides) > planned, "проба бессмысленна: колода и так не растёт"
    for cap in (planned, planned + 1):
        capped, _ = build_deck_ir(
            spec, plan, cfg.variants[1], pack=pack, max_slides=cap
        )
        assert len(capped.slides) == cap, (cap, len(capped.slides))


def test_deck_ceiling_comes_from_config():
    """Потолок приходит из `deck.max_slides` на каждом пути сборки.

    Путей два: первая сборка варианта и пересборка после исправления.
    Забыть потолок на втором — значит получить колоду в 16 слайдов после
    правки, хотя первая сборка удержалась в пятнадцати.
    """
    import inspect

    from deckwright import pipeline

    source = inspect.getsource(pipeline)
    assert source.count("build_deck_ir(") == source.count("max_slides="), (
        "у какого-то вызова вёрстки потолок не передан"
    )
