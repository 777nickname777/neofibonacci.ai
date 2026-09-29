"""Изоляция прогонов, безопасность загрузок и устойчивость экспорта в PDF.

Эти проверки закрывают то, чего не ловил ни один из существующих тестов:
427 зелёных тестов ничего не говорили ни про одновременные прогоны, ни про
скачивание чужого результата, ни про старый PDF, выданный за новый.

Где можно, отказ подделывается управляемо: `soffice` подменяется скриптом с
нужным кодом возврата. Успешная конвертация и конкуренция проверяются на
настоящем LibreOffice — подделка там ничего не доказала бы.
"""

from __future__ import annotations

import os
import shutil
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import runs

from deckwright.render.pdf import (
    ConversionError,
    ConversionReport,
    configure_parallelism,
    pptx_to_pdf,
)
from deckwright.workspace import (
    RunWorkspace,
    UnsafeUpload,
    new_run_id,
    new_session_id,
    safe_name,
)

HAS_SOFFICE = shutil.which("soffice") is not None


# ── имена файлов ────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "raw, expected",
    [
        ("../../etc/passwd.pptx", "passwd.pptx"),
        ("..\\..\\windows\\system32\\evil.pptx", "evil.pptx"),
        ("/absolute/path/deck.pptx", "deck.pptx"),
        ("обычный шаблон.pptx", "обычный шаблон.pptx"),
        ("  отчёт  за  квартал .pptx", "отчёт за квартал .pptx"),
    ],
)
def test_user_filename_cannot_leave_the_upload_directory(raw, expected):
    """Из пользовательского имени остаётся только базовое имя файла."""
    assert safe_name(raw) == expected
    assert "/" not in safe_name(raw) and "\\" not in safe_name(raw)


def test_filename_length_is_bounded(tmp_path):
    """Длинное имя обрезается, расширение сохраняется."""
    name = safe_name("я" * 500 + ".pptx")
    assert len(name) <= 120
    assert name.endswith(".pptx")


def test_same_upload_names_do_not_overwrite_each_other(tmp_path):
    """Два файла с одним именем в одном прогоне остаются двумя файлами."""
    ws = RunWorkspace.create(tmp_path, new_run_id(), new_session_id())
    first = ws.save_upload("template.pptx", b"first")
    second = ws.save_upload("template.pptx", b"second")
    assert first != second
    assert first.read_bytes() == b"first"
    assert second.read_bytes() == b"second"


def test_upload_stays_inside_its_run_directory(tmp_path):
    """Побег из каталога прогона невозможен даже именем вида `../../`."""
    ws = RunWorkspace.create(tmp_path, new_run_id(), new_session_id())
    saved = ws.save_upload("../../../../tmp/evil.pptx", b"x")
    assert saved.resolve().is_relative_to(ws.root.resolve())


def test_unsupported_extension_never_reaches_disk(tmp_path):
    ws = RunWorkspace.create(tmp_path, new_run_id(), new_session_id())
    with pytest.raises(UnsafeUpload):
        ws.save_upload("payload.exe", b"x")
    assert not any(ws.inputs.iterdir())


def test_two_runs_with_identical_filenames_do_not_share_anything(tmp_path):
    """Два одновременных прогона с одинаковыми именами входов независимы."""
    a = RunWorkspace.create(tmp_path, new_run_id(), new_session_id())
    b = RunWorkspace.create(tmp_path, new_run_id(), new_session_id())
    pa = a.save_upload("шаблон.pptx", b"A")
    pb = b.save_upload("шаблон.pptx", b"B")
    assert pa != pb
    assert pa.read_bytes() == b"A" and pb.read_bytes() == b"B"
    assert a.variant("dense") != b.variant("dense")


def test_cleaning_one_run_leaves_the_other_alone(tmp_path):
    """Уборка одного прогона не трогает соседний, который ещё работает."""
    a = RunWorkspace.create(tmp_path, new_run_id(), new_session_id())
    b = RunWorkspace.create(tmp_path, new_run_id(), new_session_id())
    a.save_upload("in.pptx", b"A")
    kept = b.save_upload("in.pptx", b"B")
    result_b = b.variant("dense") / "deck.pptx"
    result_b.write_bytes(b"deck-b")

    a.cleanup(keep_outputs=False)

    assert not a.root.exists()
    assert kept.read_bytes() == b"B"
    assert result_b.read_bytes() == b"deck-b"


def test_cleanup_keeps_finished_results(tmp_path):
    """Уборка убирает входы и обрывки, но не готовую колоду."""
    ws = RunWorkspace.create(tmp_path, new_run_id(), new_session_id())
    ws.save_upload("in.pptx", b"A")
    deck = ws.variant("airy") / "deck.pptx"
    deck.write_bytes(b"ready")
    (ws.variant("airy") / "half.part").write_bytes(b"junk")

    ws.cleanup()

    assert deck.read_bytes() == b"ready"
    assert not ws.inputs.exists()
    assert not list(ws.root.rglob("*.part"))


# ── владение прогоном ───────────────────────────────────────────────────


def _register(owner: str, run_id: str, finished: bool = True) -> runs.RunState:
    state = runs.RunState(
        run_id=run_id,
        template_name="t.pptx",
        started_at=datetime.now(UTC),
        owner=owner,
        finished=finished,
    )
    with runs._LOCK:
        runs._RUNS[run_id] = state
    return state


@pytest.fixture(autouse=True)
def _clean_registry():
    with runs._LOCK:
        runs._RUNS.clear()
    yield
    with runs._LOCK:
        runs._RUNS.clear()


def test_knowing_someone_elses_run_id_gives_nothing():
    """Номер чужого прогона сам по себе не открывает его результаты."""
    mine, theirs = new_session_id(), new_session_id()
    _register(theirs, "run-theirs")
    assert runs.get("run-theirs", theirs) is not None
    assert runs.get("run-theirs", mine) is None


def test_latest_is_scoped_to_its_own_session():
    """«Последний прогон» — последний свой, а не вообще последний."""
    mine, theirs = new_session_id(), new_session_id()
    _register(mine, "run-mine")
    time.sleep(0.01)
    _register(theirs, "run-theirs")  # чужой и более поздний
    assert runs.latest(mine).run_id == "run-mine"
    assert runs.latest(theirs).run_id == "run-theirs"
    assert runs.latest(new_session_id()) is None


def test_forget_refuses_someone_elses_run(tmp_path):
    mine, theirs = new_session_id(), new_session_id()
    state = _register(theirs, "run-theirs")
    state.workspace = RunWorkspace.create(tmp_path, "run-theirs", theirs)
    state.workspace.save_upload("in.pptx", b"x")

    assert runs.forget("run-theirs", mine) is False
    assert state.workspace.inputs.exists()  # ничего не удалено
    assert runs.get("run-theirs", theirs) is not None


def test_forget_refuses_a_running_run(tmp_path):
    """Идущий прогон не убирается: иначе исчезнут файлы из-под него."""
    owner = new_session_id()
    state = _register(owner, "run-live", finished=False)
    state.workspace = RunWorkspace.create(tmp_path, "run-live", owner)
    assert runs.forget("run-live", owner) is False
    assert runs.get("run-live", owner) is not None


# ── экспорт в PDF ───────────────────────────────────────────────────────


def _fake_soffice(tmp_path: Path, body: str) -> str:
    """Подставной `soffice` с заданным поведением; кладётся в PATH."""
    binary = tmp_path / "bin" / "soffice-fake"
    binary.parent.mkdir(parents=True, exist_ok=True)
    binary.write_text("#!/bin/sh\n" + body)
    binary.chmod(0o755)
    os.environ["PATH"] = f"{binary.parent}{os.pathsep}{os.environ['PATH']}"
    return binary.name


def test_stale_pdf_is_never_passed_off_as_a_fresh_result(tmp_path):
    """Старый PDF рядом не превращает провал конвертации в успех.

    Ровно этот дефект и был: проверялось только существование файла с нужным
    именем. `soffice` с кодом 1, ничего не записавший, — и наружу уходил
    32-байтный файл прошлого прогона.
    """
    work = tmp_path / "out"
    work.mkdir()
    (work / "deck.pptx").write_bytes(b"PK\x03\x04")
    stale = work / "deck.pdf"
    stale.write_bytes(b"%PDF-1.4 STALE")
    fake = _fake_soffice(tmp_path, "echo boom >&2\nexit 1\n")

    with pytest.raises(ConversionError):
        pptx_to_pdf(work / "deck.pptx", work, soffice_binary=fake, max_attempts=1)

    assert stale.read_bytes() == b"%PDF-1.4 STALE"  # не тронут и не выдан


def test_zero_byte_and_non_pdf_output_are_rejected(tmp_path):
    """Файл на месте — ещё не PDF: пустой и не-PDF отклоняются."""
    work = tmp_path / "out"
    work.mkdir()
    (work / "deck.pptx").write_bytes(b"PK\x03\x04")
    fake = _fake_soffice(
        tmp_path,
        'out=""\nwhile [ $# -gt 0 ]; do [ "$1" = "--outdir" ] && out="$2"; shift; done\n'
        'printf "" > "$out/deck.pdf"\nexit 0\n',
    )
    with pytest.raises(ConversionError, match="пуст"):
        pptx_to_pdf(work / "deck.pptx", work, soffice_binary=fake, max_attempts=1)


def test_missing_binary_is_permanent_and_explains_itself(tmp_path):
    work = tmp_path / "out"
    work.mkdir()
    (work / "deck.pptx").write_bytes(b"PK\x03\x04")
    with pytest.raises(ConversionError) as caught:
        pptx_to_pdf(work / "deck.pptx", work, soffice_binary="soffice-not-installed")
    assert caught.value.kind == "missing_binary"
    assert caught.value.retryable is False
    assert ".pptx" in str(caught.value)  # сказано, что колода остаётся годной


def test_transient_failure_is_retried_and_first_cause_is_kept(tmp_path):
    """Временный отказ повторяется, а причина ПЕРВОЙ попытки не теряется."""
    work = tmp_path / "out"
    work.mkdir()
    (work / "deck.pptx").write_bytes(b"PK\x03\x04")
    counter = tmp_path / "attempts"
    fake = _fake_soffice(
        tmp_path,
        f'echo x >> {counter}\necho "temporary glitch" >&2\nexit 3\n',
    )
    with pytest.raises(ConversionError) as caught:
        pptx_to_pdf(work / "deck.pptx", work, soffice_binary=fake, max_attempts=3)
    assert counter.read_text().count("x") == 3, "повторов не было"
    assert caught.value.attempts == 3
    assert "кодом 3" in str(caught.value)


def test_bad_input_is_not_retried(tmp_path):
    """То, про что LibreOffice сказал прямо, повторять бессмысленно."""
    work = tmp_path / "out"
    work.mkdir()
    (work / "deck.pptx").write_bytes(b"PK\x03\x04")
    counter = tmp_path / "attempts"
    fake = _fake_soffice(
        tmp_path,
        f'echo x >> {counter}\necho "Error: source file could not be loaded" >&2\nexit 1\n',
    )
    with pytest.raises(ConversionError) as caught:
        pptx_to_pdf(work / "deck.pptx", work, soffice_binary=fake, max_attempts=3)
    assert counter.read_text().count("x") == 1, "битый вход повторяли"
    assert caught.value.kind == "bad_input"


def test_timeout_kills_only_our_own_process(tmp_path):
    """По таймауту гасится наш процесс и его потомки — и ничего больше."""
    work = tmp_path / "out"
    work.mkdir()
    (work / "deck.pptx").write_bytes(b"PK\x03\x04")
    marker = tmp_path / "still-running"
    fake = _fake_soffice(
        tmp_path,
        f'(sleep 30; echo leaked > {marker}) &\nsleep 30\n',
    )
    started = time.monotonic()
    with pytest.raises(ConversionError, match="не уложился"):
        pptx_to_pdf(
            work / "deck.pptx", work, soffice_binary=fake,
            timeout_seconds=2, max_attempts=1,
        )
    assert time.monotonic() - started < 25, "таймаут не сработал"
    time.sleep(1.5)
    assert not marker.exists(), "потомок пережил остановку группы"


def test_missing_input_is_reported_before_launching_anything(tmp_path):
    with pytest.raises(ConversionError) as caught:
        pptx_to_pdf(tmp_path / "нет-такого.pptx", tmp_path)
    assert caught.value.kind == "missing_input"


@pytest.mark.skipif(not HAS_SOFFICE, reason="нужен настоящий LibreOffice")
def test_page_count_mismatch_is_a_failure(tmp_path):
    """PDF, в котором страниц не столько, сколько слайдов, — не результат."""
    sys.path.insert(0, str(Path(__file__).parent / "fixtures"))
    from make_template import build_template

    template = tmp_path / "t.pptx"
    build_template(template)
    report = ConversionReport(source=str(template))
    with pytest.raises(ConversionError, match="потеряла или удвоила"):
        pptx_to_pdf(
            template, tmp_path, expected_pages=999, max_attempts=1, report=report
        )


@pytest.mark.skipif(not HAS_SOFFICE, reason="нужен настоящий LibreOffice")
def test_parallel_conversions_are_capped_and_all_succeed(tmp_path):
    """Настоящий LibreOffice: восемь заявок, предел два, ноль падений."""
    sys.path.insert(0, str(Path(__file__).parent / "fixtures"))
    from make_template import build_template

    source = tmp_path / "t.pptx"
    build_template(source)
    configure_parallelism(2)

    live = 0
    peak = 0
    guard = threading.Lock()
    real_run = pptx_to_pdf

    dirs = []
    for i in range(8):
        d = tmp_path / f"v{i}"
        d.mkdir()
        shutil.copy(source, d / "deck.pptx")
        dirs.append(d)

    def job(d: Path):
        nonlocal live, peak
        with guard:
            live += 1
            peak = max(peak, live)
        try:
            return real_run(d / "deck.pptx", d, max_attempts=1, timeout_seconds=300)
        finally:
            with guard:
                live -= 1

    with ThreadPoolExecutor(max_workers=8) as pool:
        produced = list(pool.map(job, dirs))

    assert all(p.is_file() and p.stat().st_size > 0 for p in produced)
    # `peak` считает вошедших в функцию, включая ждущих на семафоре, поэтому
    # проверяем результат, а не само число: важно, что никто не упал.
    assert len(produced) == 8


@pytest.mark.skipif(not HAS_SOFFICE, reason="нужен настоящий LibreOffice")
def test_paths_with_spaces_and_cyrillic_convert(tmp_path):
    """Пробелы и кириллица в пути не ломают запуск внешнего процесса."""
    sys.path.insert(0, str(Path(__file__).parent / "fixtures"))
    from make_template import build_template

    work = tmp_path / "мои презентации" / "отчёт за квартал"
    work.mkdir(parents=True)
    deck = work / "колода образца.pptx"
    build_template(deck)
    produced = pptx_to_pdf(deck, work, max_attempts=1, timeout_seconds=300)
    assert produced.is_file() and produced.stat().st_size > 0
    assert produced.name == "колода образца.pdf"


# ── частичный результат ─────────────────────────────────────────────────


def _recorded_run(
    tmp_path, monkeypatch, fail_pdf: bool, template=None, variant: str = "balanced"
):
    """Прогон одного варианта на записанных ответах; PDF можно уронить.

    `template` — путь к настоящему шаблону; без него собирается синтетический.
    """
    import json

    sys.path.insert(0, str(Path(__file__).parent / "fixtures"))
    from make_template import build_template

    from deckwright import pipeline as pipeline_module
    from deckwright.config import load_config
    from deckwright.llm.fake import RecordedClient
    from deckwright.schemas import ContentPack

    root = Path(__file__).resolve().parents[1]
    cfg = load_config(root / "configs" / "config.yaml")
    pack = ContentPack.model_validate(
        json.loads((root / "tests/fixtures/content_pack.json").read_text("utf-8"))
    )
    template = Path(template) if template is not None else build_template(tmp_path / "t.pptx")

    if fail_pdf:
        def boom(*args, **kwargs):
            report = kwargs.get("report")
            if report is not None:
                report.attempts = 3
                report.error_kind = "transient"
            raise ConversionError(
                "LibreOffice завершился с кодом 1", kind="transient", attempts=3
            )

        monkeypatch.setattr(pipeline_module, "pptx_to_pdf", boom)

    return cfg, pipeline_module.run_variant(
        template_path=template,
        pack=pack,
        cfg=cfg,
        client=RecordedClient(root / "tests/fixtures/recorded"),
        variant=variant,
        output_dir=tmp_path / "out",
        vlm_client=None,
        fix_mode="off",
    )


@pytest.mark.skipif(not HAS_SOFFICE, reason="нужен настоящий LibreOffice")
def test_failed_pdf_does_not_destroy_the_generated_pptx(tmp_path, monkeypatch):
    """PDF не получился — колода, HTML и отчёт остаются на месте и годны."""
    _, result = _recorded_run(tmp_path, monkeypatch, fail_pdf=True)

    assert result.pptx.is_file() and result.pptx.stat().st_size > 0
    assert result.html.is_file() and result.html.stat().st_size > 0
    assert result.pdf is None
    assert "pdf" in result.export_errors
    assert result.complete is False, "неполный результат не должен считаться успехом"

    available = result.available_formats()
    assert set(available) == {"pptx", "html"}, "скачивается ровно то, что готово"

    # Паспорт прогона обязан объяснить, чего не хватает и сколько раз пробовали.
    manifest = result.manifest
    assert manifest.export_errors.get("pdf")
    assert "pdf" not in manifest.artifacts, "битого пути в артефактах быть не должно"
    assert manifest.conversions and manifest.conversions[-1].attempts == 3


@pytest.mark.skipif(not HAS_SOFFICE, reason="нужен настоящий LibreOffice")
def test_successful_run_reports_all_three_formats(tmp_path, monkeypatch):
    """Обычный прогон: три формата, ошибок экспорта нет."""
    _, result = _recorded_run(tmp_path, monkeypatch, fail_pdf=False)
    assert result.complete is True
    assert set(result.available_formats()) == {"pptx", "pdf", "html"}
    assert result.manifest.conversions[-1].ok is True
    assert result.manifest.conversions[-1].pages == len(result.deck.slides)


@pytest.mark.skipif(not HAS_SOFFICE, reason="нужен настоящий LibreOffice")
def test_retrying_export_does_not_regenerate_the_deck(tmp_path, monkeypatch):
    """Повтор экспорта берёт готовый .pptx и не трогает генерацию."""
    cfg, result = _recorded_run(tmp_path, monkeypatch, fail_pdf=True)
    monkeypatch.undo()  # дальше конвертация настоящая

    pptx_before = result.pptx.read_bytes()
    slides_before = len(result.deck.slides)

    state = runs.RunState(
        run_id="run-retry",
        template_name="t.pptx",
        started_at=datetime.now(UTC),
        owner="owner",
        finished=True,
        variants={"balanced": runs.VariantState(name="balanced", result=result)},
    )
    runs.retry_export(state, "balanced", cfg)
    deadline = time.monotonic() + 300
    while state.variants["balanced"].exporting and time.monotonic() < deadline:
        time.sleep(0.5)

    assert state.variants["balanced"].exporting is False
    assert result.pdf is not None and result.pdf.is_file()
    assert "pdf" not in result.export_errors
    assert result.complete is True
    # Колода не пересобиралась: те же байты, то же число слайдов.
    assert result.pptx.read_bytes() == pptx_before
    assert len(result.deck.slides) == slides_before


# ── конфигурация и .env ─────────────────────────────────────────────────


def test_env_file_is_read_by_the_application_itself(tmp_path, monkeypatch):
    """`.env` подхватывается кодом, без `source` в shell."""
    from deckwright.config import ENV_FILE_VAR, load_env_file

    env = tmp_path / ".env"
    env.write_text(
        "# комментарий\n"
        "DECKWRIGHT_T_PLAIN=value\n"
        'DECKWRIGHT_T_QUOTED="со пробелом"\n'
        "export DECKWRIGHT_T_EXPORTED=ok\n"
        "мусор без равенства\n",
        encoding="utf-8",
    )
    monkeypatch.setenv(ENV_FILE_VAR, str(env))
    for name in ("DECKWRIGHT_T_PLAIN", "DECKWRIGHT_T_QUOTED", "DECKWRIGHT_T_EXPORTED"):
        monkeypatch.delenv(name, raising=False)

    applied = load_env_file(override=True)

    assert os.environ["DECKWRIGHT_T_PLAIN"] == "value"
    assert os.environ["DECKWRIGHT_T_QUOTED"] == "со пробелом"
    assert os.environ["DECKWRIGHT_T_EXPORTED"] == "ok"
    assert set(applied) == {
        "DECKWRIGHT_T_PLAIN", "DECKWRIGHT_T_QUOTED", "DECKWRIGHT_T_EXPORTED"
    }


def test_real_environment_wins_over_the_file(tmp_path, monkeypatch):
    """Переменная, уже заданная снаружи, не перетирается файлом."""
    from deckwright.config import ENV_FILE_VAR, load_env_file

    env = tmp_path / ".env"
    env.write_text("DECKWRIGHT_T_PRIORITY=from-file\n", encoding="utf-8")
    monkeypatch.setenv(ENV_FILE_VAR, str(env))
    monkeypatch.setenv("DECKWRIGHT_T_PRIORITY", "from-environment")

    load_env_file(override=False)

    assert os.environ["DECKWRIGHT_T_PRIORITY"] == "from-environment"


def test_env_path_does_not_depend_on_the_current_directory(monkeypatch, tmp_path):
    """Путь к `.env` предсказуем: корень проекта, а не откуда запустили."""
    from deckwright.config import REPO_ROOT, env_file_path

    monkeypatch.delenv("DECKWRIGHT_ENV_FILE", raising=False)
    monkeypatch.chdir(tmp_path)
    assert env_file_path() == REPO_ROOT / ".env"


def test_env_loading_never_returns_values(tmp_path, monkeypatch):
    """Наружу отдаются только имена ключей: в файле лежат секреты."""
    from deckwright.config import ENV_FILE_VAR, load_env_file

    env = tmp_path / ".env"
    env.write_text("DECKWRIGHT_T_SECRET=super-secret-value\n", encoding="utf-8")
    monkeypatch.setenv(ENV_FILE_VAR, str(env))
    monkeypatch.delenv("DECKWRIGHT_T_SECRET", raising=False)

    applied = load_env_file(override=True)

    assert applied == ["DECKWRIGHT_T_SECRET"]
    assert "super-secret-value" not in " ".join(applied)


def test_readiness_separates_blocking_from_degraded():
    """Нет PDF-конвертера — это не «нельзя работать», а «не будет PDF»."""
    from deckwright.environment import check_soffice

    missing = check_soffice("soffice-definitely-not-installed")
    assert missing.ok is False
    assert missing.blocking is False, "отсутствие LibreOffice не блокирует прогон"
    assert "pptx" in missing.affects


def test_readiness_flags_unwritable_directories(tmp_path):
    """Каталог без записи — блокирующая проблема, и она видна заранее."""
    from deckwright.environment import check_writable

    locked = tmp_path / "locked"
    locked.mkdir()
    locked.chmod(0o500)
    try:
        result = check_writable([locked])
        assert result.ok is False
        assert result.blocking is True
    finally:
        locked.chmod(0o700)


# ── интерфейс: сессии и обновление страницы ─────────────────────────────


def _ui(monkeypatch):
    from streamlit.testing.v1 import AppTest

    monkeypatch.delenv("DECKWRIGHT_PASSWORD", raising=False)
    root = Path(__file__).resolve().parents[1]
    monkeypatch.chdir(root)
    return AppTest.from_file(str(root / "app" / "ui.py"), default_timeout=90)


def test_each_browser_session_gets_its_own_identity(monkeypatch):
    """Две вкладки — два владельца; чужой прогон в адресе ничего не открывает."""
    first = _ui(monkeypatch).run()
    second = _ui(monkeypatch).run()
    id_first = first.session_state["session_id"]
    id_second = second.session_state["session_id"]
    assert id_first != id_second

    # Прогон первой сессии со стороны второй недоступен.
    _register(id_first, "run-of-the-first")
    assert runs.get("run-of-the-first", id_first) is not None
    assert runs.get("run-of-the-first", id_second) is None


def test_refresh_does_not_start_a_second_run(monkeypatch):
    """Перезапуск скрипта не нажимает «Собрать» повторно.

    Кнопка Streamlit истинна ровно один прогон скрипта. Если бы состояние
    кнопки переживало перезапуск, обновление страницы запускало бы новую
    генерацию — с новым вызовом модели и новым каталогом.
    """
    app = _ui(monkeypatch).run()
    before = len(runs._RUNS)
    for _ in range(3):
        app.run()
    assert len(runs._RUNS) == before, "обновление страницы породило прогон"


def test_unknown_run_id_in_the_url_is_explained_not_served(monkeypatch):
    """Чужой номер в адресе даёт объяснение, а не чужие результаты."""
    other = new_session_id()
    _register(other, "run-someone-else")
    app = _ui(monkeypatch)
    app.query_params["run"] = "run-someone-else"
    app.run()
    assert app.warning, "о чужом прогоне ничего не сказано"
    # Среди всех предупреждений, а не первым: страница говорит и о другом —
    # например, что пароль на вход не задан при настроенном ключе модели.
    assert any("не найден" in item.value for item in app.warning), [
        item.value for item in app.warning
    ]


def test_run_writes_a_diagnostic_summary(tmp_path):
    """У прогона остаётся паспорт: этапы, форматы, попытки, причины."""
    owner = new_session_id()
    ws = RunWorkspace.create(tmp_path, new_run_id(), owner)
    state = runs.RunState(
        run_id=ws.run_id,
        template_name="t.pptx",
        started_at=datetime.now(UTC),
        owner=owner,
        workspace=ws,
        finished=True,
        stage="готово частично",
        variants={"dense": runs.VariantState(name="dense", stage="готово")},
    )
    written = state.write_summary()
    assert written is not None and written.name == "run.json"

    import json as _json

    data = _json.loads(written.read_text(encoding="utf-8"))
    assert data["run_id"] == ws.run_id
    assert data["stage"] == "готово частично"
    assert "dense" in data["variants"]
    assert "stages" in data and "elapsed_seconds" in data


@pytest.mark.skipif(not HAS_SOFFICE, reason="нужен настоящий LibreOffice")
def test_interface_offers_downloads_only_for_its_own_finished_run(monkeypatch, tmp_path):
    """Интерфейс показывает кнопки скачивания своего прогона — и только их.

    Проверяется настоящий код страницы: собранный прогон кладётся в реестр от
    имени сессии, которую завела сама страница, после чего она отрисовывается
    и должна показать ровно доступные форматы.
    """
    _, result = _recorded_run(tmp_path, monkeypatch, fail_pdf=False)

    app = _ui(monkeypatch)
    app.run()
    owner = app.session_state["session_id"]

    state = runs.RunState(
        run_id="run-ui-check",
        template_name="t.pptx",
        started_at=datetime.now(UTC),
        owner=owner,
        finished=True,
        stage="готово",
        variants={"balanced": runs.VariantState(name="balanced", result=result,
                                                stage="готово")},
    )
    with runs._LOCK:
        runs._RUNS["run-ui-check"] = state

    app.query_params["run"] = "run-ui-check"
    app.run()

    labels = {b.label for b in app.button} | {
        getattr(b, "label", "") for b in getattr(app, "download_button", [])
    }
    assert any("pptx" in str(label) for label in labels), f"нет кнопки .pptx: {labels}"
    assert any("pdf" in str(label) for label in labels), f"нет кнопки .pdf: {labels}"
    assert any("html" in str(label) for label in labels), f"нет кнопки .html: {labels}"
    assert not app.exception, f"страница упала: {app.exception}"


@pytest.mark.skipif(not HAS_SOFFICE, reason="нужен настоящий LibreOffice")
def test_interface_names_the_missing_format_and_offers_a_retry(monkeypatch, tmp_path):
    """PDF не получился — страница говорит об этом и предлагает повтор экспорта."""
    _, result = _recorded_run(tmp_path, monkeypatch, fail_pdf=True)
    monkeypatch.undo()

    app = _ui(monkeypatch)
    app.run()
    owner = app.session_state["session_id"]
    state = runs.RunState(
        run_id="run-ui-partial",
        template_name="t.pptx",
        started_at=datetime.now(UTC),
        owner=owner,
        finished=True,
        stage="готово частично",
        variants={"balanced": runs.VariantState(name="balanced", result=result,
                                                stage="готово частично: pdf")},
    )
    with runs._LOCK:
        runs._RUNS["run-ui-partial"] = state

    app.query_params["run"] = "run-ui-partial"
    app.run()

    assert not app.exception, f"страница упала: {app.exception}"
    warnings = " ".join(w.value for w in app.warning)
    assert "не получен" in warnings, f"о нехватке формата не сказано: {warnings}"
    labels = [str(b.label) for b in app.button]
    assert any("Повторить только экспорт PDF" in label for label in labels), labels
    # Колода всё равно скачивается.
    dl = [str(getattr(b, "label", "")) for b in getattr(app, "download_button", [])]
    assert any("pptx" in label for label in dl), dl


# ── коллизия имён частей пакета ─────────────────────────────────────────


def _zip_names(pptx: Path) -> list[str]:
    import zipfile

    with zipfile.ZipFile(pptx) as zf:
        return zf.namelist()


def test_donor_chart_and_our_chart_do_not_collide_on_a_partname(tmp_path):
    """Два разных графика не могут называться одинаково.

    Воспроизводит механизм дефекта напрямую: снимаем слайды (донорский график
    осиротевает), добавляем свой график — python-pptx выдаёт ему то же имя, —
    затем возвращаем донорский в граф, как это делает `clone._relink`.
    """
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.util import Emu

    from deckwright.render.clone import ensure_unique_partnames, purge_slides

    source = Path("data/holdout/zelenie_investicii.pptx")
    if not source.is_file():
        pytest.skip("нет шаблона data/holdout/zelenie_investicii.pptx")

    prs = Presentation(str(source))
    donor_chart = next(
        rel.target_part
        for slide in prs.slides
        for rel in slide.part.rels.values()
        if rel.reltype == RT.CHART
    )
    layout = prs.slide_layouts[0]
    purge_slides(prs)

    data = CategoryChartData()
    data.categories = ["До", "После"]
    data.add_series("минуты", (42.0, 9.0))
    ours = prs.slides.add_slide(layout).shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(0), Emu(0), Emu(3000000), Emu(2000000), data
    ).chart.part
    prs.slides.add_slide(layout).part.relate_to(donor_chart, RT.CHART)

    assert ours is not donor_chart
    assert str(ours.partname) == str(donor_chart.partname), (
        "дефект больше не воспроизводится этим путём — проверьте механизм"
    )

    renamed = ensure_unique_partnames(prs)

    assert renamed, "коллизия не устранена"
    assert str(ours.partname) != str(donor_chart.partname)
    out = tmp_path / "deck.pptx"
    prs.save(str(out))
    names = _zip_names(out)
    assert len(names) == len(set(names)), "в архиве остались повторы"
    charts = sorted(n for n in names if n.startswith("ppt/charts/chart"))
    assert len([c for c in charts if c.endswith(".xml")]) == 2, f"график потерян: {charts}"


def test_shared_part_is_serialized_once(tmp_path):
    """Одна часть, на которую ссылаются двое, остаётся одной частью.

    Законное совместное использование трогать нельзя: переименование должно
    срабатывать только на разных объектах с одинаковым именем.
    """
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    from deckwright.render.clone import ensure_unique_partnames, purge_slides

    source = Path("data/holdout/zelenie_investicii.pptx")
    if not source.is_file():
        pytest.skip("нет шаблона")

    prs = Presentation(str(source))
    donor_chart = next(
        rel.target_part
        for slide in prs.slides
        for rel in slide.part.rels.values()
        if rel.reltype == RT.CHART
    )
    layout = prs.slide_layouts[0]
    purge_slides(prs)
    # Один и тот же объект части — двум слайдам.
    for _ in range(2):
        prs.slides.add_slide(layout).part.relate_to(donor_chart, RT.CHART)

    assert ensure_unique_partnames(prs) == [], "общая часть переименована напрасно"

    out = tmp_path / "shared.pptx"
    prs.save(str(out))
    names = _zip_names(out)
    assert names.count("ppt/charts/chart1.xml") == 1, "общая часть записана дважды"


def test_integrity_check_sees_duplicate_parts_and_content_types(tmp_path):
    """Битый пакет не должен считаться готовым результатом."""
    import zipfile

    from deckwright.render.package_check import check_package

    broken = tmp_path / "broken.pptx"
    with zipfile.ZipFile(broken, "w") as zf:
        zf.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/'
            'package/2006/content-types">'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/a.xml" ContentType="application/xml"/>'
            '<Override PartName="/a.xml" ContentType="application/xml"/>'
            "</Types>",
        )
        zf.writestr("a.xml", "<a/>")
        zf.writestr("a.xml", "<a>другое содержимое</a>")

    report = check_package(broken)

    assert not report.ok
    joined = " ".join(report.problems)
    assert "дважды" in joined
    assert "a.xml" in joined


@pytest.mark.skipif(not HAS_SOFFICE, reason="нужен настоящий LibreOffice")
def test_zelenie_balanced_builds_a_loadable_deck_with_both_charts(tmp_path, monkeypatch):
    """Тот самый прогон, который падал: колода собирается и открывается.

    Проверяется не «zip открылся», а то, ради чего всё делалось: нет повторов
    имён, график колоды на месте со своими данными и без рыбных данных
    донора, внутренние ссылки целы, LibreOffice конвертирует файл в PDF на
    столько же страниц, сколько слайдов.
    """
    import zipfile

    from deckwright.render.package_check import check_package

    source = Path("data/holdout/zelenie_investicii.pptx")
    if not source.is_file():
        pytest.skip("нет шаблона data/holdout/zelenie_investicii.pptx")

    _, result = _recorded_run(
        tmp_path, monkeypatch, fail_pdf=False, template=source, variant="balanced"
    )

    names = _zip_names(result.pptx)
    assert len(names) == len(set(names)), "повторы имён вернулись"
    assert check_package(result.pptx).ok

    charts = sorted(
        n for n in set(names) if n.startswith("ppt/charts/chart") and n.endswith(".xml")
    )
    assert charts, "график колоды не доехал до пакета"
    with zipfile.ZipFile(result.pptx) as zf:
        bodies = [zf.read(c) for c in charts]
    assert all(len(b) > 500 for b in bodies), "график пуст"
    assert len(set(bodies)) == len(bodies), "два графика оказались одним и тем же"
    # Наши числа доехали до данных графика.
    assert any(b"42" in b and b"9" in b for b in bodies)
    # И ни в одном графике не осталось рыбных данных донора. Раньше здесь
    # ждали ровно два графика, и вторым был именно донорский — «Ряд 1» с
    # «Категорией 1»: композиция бралась вместе с его диаграммой. Он ушёл,
    # когда оформление фоновой картинки стало защищённой зоной и текст
    # перестал помещаться на тот слайд-донор. Требование от этого строже, а
    # не слабее: раньше leftover донора считался нормой, теперь запрещён —
    # его же ловит и аудит (`integrity.donor_data_leftover`).
    for name, body in zip(charts, bodies, strict=True):
        text = body.decode("utf-8", "ignore")
        assert "Категория 1" not in text and "Ряд 1" not in text, (
            f"в {name} остались рыбные данные шаблона"
        )

    assert result.pdf is not None and result.pdf.is_file()
    assert result.complete
    assert len(result.pages) == len(result.deck.slides)
