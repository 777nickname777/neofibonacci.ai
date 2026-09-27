"""Шаблоны без плейсхолдеров: композиции разбираются по фигурам.

Шаблон-бланк «Дорожная карта» (десять слайдов, один layout, ни одного
плейсхолдера, всё — обычные фигуры) сервис раньше собирал на обложке и финале:
у рабочих рамок нулевые поля и высота ровно в строку, и фиттер считал, что
туда не влезает ни строки. Маршрут, схема, таблица и карточки ролей терялись,
в заголовки ехали списки.

Проверяется на двух шаблонах, чтобы не подогнать под один файл: настоящий
бланк с говорящими именами фигур и синтетический другой структуры без имён.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest
from pptx import Presentation

from deckwright.config import load_config
from deckwright.llm.fake import RecordedClient
from deckwright.parse.opener import parse_template
from deckwright.pipeline import run_variant
from deckwright.schemas import ContentPack, PatternClass, SlotRole

sys.path.insert(0, str(Path(__file__).parent / "fixtures"))
from make_blank_template import PLACEHOLDER_TEXTS, build

ROOT = Path(__file__).resolve().parents[1]
CONFIG = ROOT / "configs" / "config.yaml"
ROADMAP = ROOT / "data" / "holdout" / "dorozhnaya_karta.pptx"
FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="module")
def roadmap():
    if not ROADMAP.exists():
        pytest.skip("нет шаблона data/holdout/dorozhnaya_karta.pptx")
    return parse_template(ROADMAP)


@pytest.fixture(scope="module")
def blank_path(tmp_path_factory) -> Path:
    return build(tmp_path_factory.mktemp("blank") / "blank.pptx")


@pytest.fixture(scope="module")
def blank(blank_path):
    return parse_template(blank_path)


@pytest.fixture(scope="module")
def pack() -> ContentPack:
    return ContentPack.model_validate(
        json.loads((FIXTURES / "content_pack.json").read_text("utf-8"))
    )


def _roles(pattern) -> set[SlotRole]:
    return {slot.role for slot in pattern.slots}


def _template_texts(path: Path) -> set[str]:
    """Тексты надписей шаблона, кроме чисел: номера пишет сама колода."""
    texts = set()
    for slide in Presentation(str(path)).slides:
        for shape in slide.shapes:
            text = shape.text_frame.text.strip() if shape.has_text_frame else ""
            if text and not text.isdigit():
                texts.add(text)
    return texts


def _deck(template: Path, pack, tmp_path: Path, variant: str):
    return run_variant(
        template_path=template,
        pack=pack,
        cfg=load_config(CONFIG),
        client=RecordedClient(FIXTURES / "recorded"),
        variant=variant,
        output_dir=tmp_path / variant,
    )


# ── Разбор ───────────────────────────────────────────────────────────────────


def test_every_roadmap_slide_becomes_a_composition(roadmap):
    """Десять слайдов — десять композиций; обложка и финал опознаны по кеглю."""
    assert sorted(p.donor_slide_index for p in roadmap.patterns) == list(range(1, 11))
    by_slide = {p.donor_slide_index: p for p in roadmap.patterns}
    assert by_slide[1].pattern_class is PatternClass.TITLE
    assert by_slide[10].pattern_class is PatternClass.CLOSING
    assert roadmap.cover_pattern_id == by_slide[1].id
    assert roadmap.closing_pattern_id == by_slide[10].id


def test_roadmap_roles_come_from_geometry(roadmap):
    """Заголовок, подзаголовок под ним и номер страницы — на каждом рабочем слайде."""
    for pattern in roadmap.patterns:
        if pattern.donor_slide_index in (1, 10):
            continue
        roles = _roles(pattern)
        assert {SlotRole.TITLE, SlotRole.SUBTITLE, SlotRole.SLIDE_NUMBER} <= roles, (
            pattern.id,
            roles,
        )
        # Номер страницы — не место под содержание.
        assert not any(
            slot.role is SlotRole.KPI_LABEL and (slot.placeholder_text or "").isdigit()
            for slot in pattern.slots
        )


def test_roadmap_repeaters_and_their_numbers(roadmap):
    """Маршрут, строки плана, роли, ресурсы, риски, проверка — повторители.

    Номера «01…06» в узлах — порядковые номера элементов, а не места.
    """
    by_slide = {p.donor_slide_index: p for p in roadmap.patterns}
    for index in range(2, 11):
        assert by_slide[index].repeaters, f"на слайде {index} не найден повторитель"
    route = by_slide[2].repeaters[0]
    assert route.observed_count == 6
    assert any(slot.role is SlotRole.ORDINAL for slot in route.item_slots)


def test_roadmap_text_frames_hold_text(roadmap):
    """Рамка ровно в строку с нулевыми полями вмещает строку своего кегля."""
    from deckwright.layout.fitter import capacity_lines
    from deckwright.layout.text_metrics import metrics_for_spec

    source = metrics_for_spec(roadmap)
    metrics = getattr(source, "metrics", source)
    title = next(s for s in roadmap.patterns[1].slots if s.role is SlotRole.TITLE)
    assert title.box.inset_y == 0
    assert capacity_lines(metrics, title.style.size_pt, title.box) >= 1


def test_blank_template_without_names_parses_the_same_way(blank):
    """Другая структура, имена фигур стёрты — роли и повторители те же."""
    assert blank.cover_pattern_id is not None and blank.closing_pattern_id is not None
    content = [p for p in blank.patterns if p.pattern_class not in (
        PatternClass.TITLE, PatternClass.CLOSING)]
    assert len(content) == 2
    for pattern in content:
        assert {SlotRole.TITLE, SlotRole.SUBTITLE, SlotRole.SLIDE_NUMBER} <= _roles(pattern)
        assert pattern.repeaters and pattern.repeaters[0].observed_count == 4
    steps = next(p for p in content if p.repeaters[0].axis == "vertical")
    roles = {slot.role for slot in steps.repeaters[0].item_slots}
    assert SlotRole.ORDINAL in roles
    # Описание шага выросло на линию для записи под ним.
    description = max(steps.repeaters[0].item_slots, key=lambda slot: slot.box.h)
    assert description.box.h > 0.5 * 914_400


# ── Колода ───────────────────────────────────────────────────────────────────


@pytest.mark.parametrize("variant", ["dense", "balanced", "airy"])
def test_roadmap_deck_uses_its_compositions(roadmap, pack, tmp_path, variant):
    """Колода берёт рабочие композиции шаблона, а не только обложку и финал."""
    result = _deck(ROADMAP, pack, tmp_path, variant)
    content = {
        slide.pattern_id
        for slide in result.deck.slides[1:-1]
    } - {roadmap.cover_pattern_id, roadmap.closing_pattern_id}
    assert len(content) >= 3, content
    assert result.deck.slides[0].pattern_id == roadmap.cover_pattern_id


@pytest.mark.parametrize(
    "template", ["roadmap", "blank"], ids=["dorozhnaya_karta", "synthetic"]
)
def test_deck_has_no_template_text_and_numbers_its_pages(
    template, pack, tmp_path, blank_path, request
):
    """Ни одной надписи шаблона в колоде; номер страницы — номер слайда."""
    path = ROADMAP if template == "roadmap" else blank_path
    if template == "roadmap":
        request.getfixturevalue("roadmap")
    spec = parse_template(path)
    placeholders = _template_texts(path)
    if template == "blank":
        assert set(PLACEHOLDER_TEXTS) <= placeholders
    result = _deck(path, pack, tmp_path, "balanced")
    patterns = {pattern.id: pattern for pattern in spec.patterns}
    deck = Presentation(str(result.pptx))
    pairs = zip(deck.slides, result.deck.slides, strict=True)
    for index, (slide, slide_ir) in enumerate(pairs, start=1):
        texts = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        leftovers = [text for text in texts if text in placeholders]
        assert not leftovers, (index, leftovers)
        pattern = patterns[slide_ir.pattern_id]
        numbered = [s for s in pattern.slots if s.role is SlotRole.SLIDE_NUMBER]
        if numbered:
            width = len((numbered[0].placeholder_text or "").strip())
            assert str(index).zfill(width) in texts, (index, texts)


def test_progress_bar_stays_on_every_roadmap_slide(roadmap, pack, tmp_path):
    """Прогресс-бар шаблона — на каждом рабочем слайде колоды, целиком."""
    template = Presentation(str(ROADMAP))
    tops = {shape.top for shape in template.slides[1].shapes if shape.name.startswith("progress")}
    assert len(tops) == 1
    top = tops.pop()
    segments = sum(
        1 for shape in template.slides[1].shapes if shape.name.startswith("progress")
    )
    result = _deck(ROADMAP, pack, tmp_path, "dense")
    deck = Presentation(str(result.pptx))
    pairs = zip(deck.slides, result.deck.slides, strict=True)
    for index, (slide, slide_ir) in enumerate(pairs, start=1):
        if slide_ir.pattern_id in (roadmap.cover_pattern_id, roadmap.closing_pattern_id):
            continue
        found = sum(1 for shape in slide.shapes if shape.top == top and shape.height < 91_440)
        assert found == segments, (index, found)
