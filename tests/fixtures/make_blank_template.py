"""Синтетический шаблон без плейсхолдеров и без имён фигур.

Проверка разбора по фигурам не должна держаться на одном файле: настоящий
бланк `data/holdout/dorozhnaya_karta.pptx` называет фигуры (`slide-title`,
`goal-node-0`), и разбор, выучивший эти имена, прошёл бы тест на нём и сломался
на следующем. Здесь структура другая, а имена стёрты:

* номер страницы — внизу слева, а не справа;
* карточки идут строками (номер в кружке, заголовок, описание), а не
  колонками;
* сетка 2 × 2 из залитых подложек с текстом внутри;
* обложка и финал набраны крупнее рабочего заголовка;
* под описаниями — линии для записи.

Текст надписей — заглушки, которых в готовой колоде быть не должно.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Emu, Pt

INCH = 914_400
ACCENT = RGBColor(0x2B, 0x6C, 0xB0)
INK = RGBColor(0x1F, 0x29, 0x33)
MUTED = RGBColor(0x6B, 0x72, 0x80)
PANEL = RGBColor(0xEE, 0xF2, 0xF7)

# Тексты-заглушки: тест ищет их в колоде.
PLACEHOLDER_TEXTS = (
    "Название доклада",
    "Подзаголовок доклада",
    "Заголовок слайда",
    "Пояснение к слайду",
    "Шаг",
    "Опишите шаг",
    "Направление",
    "Опишите направление",
    "Спасибо",
    "Контакты для связи",
)


def _text(slide, x, y, w, h, text, size, *, bold=False, color=INK):
    box = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    frame = box.text_frame
    frame.word_wrap = True
    body = frame._txBody.bodyPr
    for side in ("lIns", "tIns", "rIns", "bIns"):
        body.set(side, "0")
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _shape(slide, kind, x, y, w, h, fill=None, line=None):
    shape = slide.shapes.add_shape(kind, Emu(x), Emu(y), Emu(w), Emu(h))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    return shape


def _writing_line(slide, x, y, w):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Emu(x), Emu(y), Emu(x + w), Emu(y)
    )
    connector.line.color.rgb = MUTED
    return connector


def _header(slide, number: int) -> None:
    _text(slide, int(0.6 * INCH), int(0.4 * INCH), int(11.5 * INCH), int(0.55 * INCH),
          "Заголовок слайда", 30, bold=True)
    _text(slide, int(0.6 * INCH), int(1.0 * INCH), int(9 * INCH), int(0.3 * INCH),
          "Пояснение к слайду", 14, color=MUTED)
    _text(slide, int(0.6 * INCH), int(7.0 * INCH), int(0.5 * INCH), int(0.25 * INCH),
          f"{number}", 11, color=MUTED)


def build(path: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Emu(int(13.333 * INCH))
    prs.slide_height = Emu(int(7.5 * INCH))
    blank = prs.slide_layouts[6]

    # 1. Обложка: заголовок крупнее рабочего.
    cover = prs.slides.add_slide(blank)
    _text(cover, int(0.8 * INCH), int(2.4 * INCH), int(9 * INCH), int(1.2 * INCH),
          "Название доклада", 48, bold=True)
    _text(cover, int(0.8 * INCH), int(3.8 * INCH), int(8 * INCH), int(0.5 * INCH),
          "Подзаголовок доклада", 18, color=MUTED)

    # 2. Шаги строками: кружок с номером, заголовок, описание, линия для записи.
    steps = prs.slides.add_slide(blank)
    _header(steps, 2)
    for index in range(4):
        top = int((1.7 + 1.25 * index) * INCH)
        _shape(steps, MSO_SHAPE.OVAL, int(0.7 * INCH), top, int(0.5 * INCH), int(0.5 * INCH),
               line=ACCENT)
        _text(steps, int(0.7 * INCH), top + int(0.12 * INCH), int(0.5 * INCH),
              int(0.25 * INCH), f"{index + 1}", 12, bold=True, color=ACCENT)
        _text(steps, int(1.5 * INCH), top, int(3 * INCH), int(0.35 * INCH), "Шаг", 18, bold=True)
        _text(steps, int(1.5 * INCH), top + int(0.42 * INCH), int(9 * INCH),
              int(0.28 * INCH), "Опишите шаг", 14, color=MUTED)
        _writing_line(steps, int(1.5 * INCH), top + int(0.95 * INCH), int(9 * INCH))

    # 3. Сетка 2 × 2 из подложек.
    grid = prs.slides.add_slide(blank)
    _header(grid, 3)
    for row in range(2):
        for column in range(2):
            x = int((0.7 + 6.1 * column) * INCH)
            y = int((1.7 + 2.6 * row) * INCH)
            _shape(grid, MSO_SHAPE.RECTANGLE, x, y, int(5.8 * INCH), int(2.3 * INCH), fill=PANEL)
            _text(grid, x + int(0.3 * INCH), y + int(0.3 * INCH), int(5.2 * INCH),
                  int(0.4 * INCH), "Направление", 20, bold=True)
            _text(grid, x + int(0.3 * INCH), y + int(0.85 * INCH), int(5.2 * INCH),
                  int(0.3 * INCH), "Опишите направление", 14, color=MUTED)

    # 4. Финал: заголовок крупнее рабочего.
    closing = prs.slides.add_slide(blank)
    _text(closing, int(0.8 * INCH), int(2.6 * INCH), int(9 * INCH), int(1.0 * INCH),
          "Спасибо", 40, bold=True)
    _text(closing, int(0.8 * INCH), int(3.8 * INCH), int(8 * INCH), int(0.4 * INCH),
          "Контакты для связи", 16, color=MUTED)

    # Имён у фигур нет: разбор не должен на них опираться.
    for slide in prs.slides:
        for shape in slide.shapes:
            shape.name = ""
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


if __name__ == "__main__":
    import sys

    print(build(Path(sys.argv[1] if len(sys.argv) > 1 else "blank_template.pptx")))
