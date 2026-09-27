"""Синтетический шаблон с фото-иллюстрациями, логотипом, значками и декором.

Шаблон «План подготовки к экзаменам» несёт на слайдах-примерах фото (студент
за компьютером, часы), и они доезжали до колоды о продажах кофеен. Сам этот
шаблон в репозиторий не положен, поэтому проверка идёт на синтетическом с той
же структурой, чтобы не подогнать разбор под один файл:

* обложка и финал — крупный плоский декор на прозрачном фоне и логотип;
* три композиции с фото: фото слева и текст справа, баннер-фото сверху и
  карточки, текст слева и фото справа;
* две композиции без фото: три карточки со значками, заголовок и абзац.

Фото здесь — сгенерированный шум с плавными переходами: непрозрачный,
многоцветный, без поля одного цвета. Декор, логотип и значки — плоские
фигуры в пару цветов на прозрачном фоне.

Запуск: `python tests/fixtures/make_photo_template.py ПУТЬ.pptx`
"""

from __future__ import annotations

import io
import random
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

INCH = 914_400
W, H = 12 * INCH, int(6.75 * INCH)
INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0xF2, 0xB7, 0x05)
CARD = RGBColor(0xFF, 0xFF, 0xFF)

PLACEHOLDER_TEXTS = (
    "План подготовки",
    "Название раздела",
    "Опишите, что нужно сделать",
    "Шаг",
    "Опишите шаг",
    "Удачи на экзамене",
)


def _photo(seed: int, size: tuple[int, int]) -> bytes:
    """Картинка «как фото»: плавный многоцветный шум без поля одного цвета."""
    rng = random.Random(seed)
    small = Image.new("RGB", (24, 16))
    small.putdata(
        [
            (rng.randrange(40, 230), rng.randrange(40, 230), rng.randrange(40, 230))
            for _ in range(24 * 16)
        ]
    )
    image = small.resize(size, Image.BICUBIC).filter(ImageFilter.GaussianBlur(2))
    grain = Image.effect_noise(size, 18).convert("RGB")
    image = Image.blend(image, grain, 0.15)
    out = io.BytesIO()
    image.save(out, "JPEG", quality=85)
    return out.getvalue()


def _flat(size: tuple[int, int], shape: str) -> bytes:
    """Плоская графика в пару цветов на прозрачном фоне: декор, логотип, значок."""
    image = Image.new("RGBA", size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    w, h = size
    if shape == "zigzag":
        draw.polygon([(0, h), (w * 0.3, 0), (w * 0.5, h * 0.6), (w * 0.8, 0), (w, h)],
                     fill=(242, 183, 5, 255))
    elif shape == "logo":
        draw.ellipse([0, 0, w - 1, h - 1], fill=(26, 26, 26, 255))
        draw.rectangle([w * 0.3, h * 0.3, w * 0.7, h * 0.7], fill=(242, 183, 5, 255))
    else:
        draw.rounded_rectangle([w * 0.1, h * 0.1, w * 0.9, h * 0.9], radius=w // 5,
                               fill=(242, 183, 5, 255))
    out = io.BytesIO()
    image.save(out, "PNG")
    return out.getvalue()


def _text(slide, x, y, w, h, text, size, *, bold=False, color=INK):
    box = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    frame = box.text_frame
    frame.word_wrap = True
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _list(slide, x, y, w, h, size):
    """Тело слайда — список из трёх пунктов-заглушек."""
    box = _text(slide, x, y, w, h, "Опишите, что нужно сделать", size)
    for _ in range(2):
        paragraph = box.text_frame.add_paragraph()
        run = paragraph.add_run()
        run.text = "Опишите шаг"
        run.font.size = Pt(size)
        run.font.color.rgb = INK
    return box


def _picture(slide, blob: bytes, x, y, w, h):
    return slide.shapes.add_picture(io.BytesIO(blob), Emu(x), Emu(y), Emu(w), Emu(h))


def _card(slide, x, y, w, h):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.line.color.rgb = ACCENT
    return card


def build(path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    blank = prs.slide_layouts[6]
    logo = _flat((200, 200), "logo")
    icon = _flat((160, 160), "icon")

    # 1. Обложка: крупный декор и логотип.
    slide = prs.slides.add_slide(blank)
    _picture(slide, _flat((800, 600), "zigzag"), 7 * INCH, int(2.5 * INCH), 5 * INCH, 4 * INCH)
    _picture(slide, logo, int(0.5 * INCH), int(0.4 * INCH), int(0.6 * INCH), int(0.6 * INCH))
    _text(slide, int(0.6 * INCH), int(2.2 * INCH), 6 * INCH, int(1.6 * INCH),
          "План подготовки", 40, bold=True)

    # 2. Фото слева, текст справа.
    slide = prs.slides.add_slide(blank)
    _picture(slide, _photo(1, (640, 760)), 0, 0, int(4.5 * INCH), H)
    _text(slide, 5 * INCH, int(0.6 * INCH), int(6.4 * INCH), int(0.9 * INCH),
          "Название раздела", 28, bold=True)
    _list(slide, 5 * INCH, int(3.1 * INCH), int(6.4 * INCH), int(3.1 * INCH), 16)

    # 3. Баннер-фото сверху и три карточки.
    slide = prs.slides.add_slide(blank)
    _picture(slide, _photo(2, (1200, 300)), 0, 0, W, int(2.2 * INCH))
    _text(slide, int(0.6 * INCH), int(2.4 * INCH), int(10.8 * INCH), int(0.8 * INCH),
          "Название раздела", 28, bold=True)
    for index in range(3):
        x = int(0.6 * INCH) + index * int(3.7 * INCH)
        _card(slide, x, int(3.5 * INCH), int(3.4 * INCH), int(2.6 * INCH))
        _text(slide, x + int(0.2 * INCH), int(3.7 * INCH), 3 * INCH, int(2.2 * INCH),
              "Опишите шаг", 14)

    # 4. Без фото: три карточки со значками.
    slide = prs.slides.add_slide(blank)
    _text(slide, int(0.6 * INCH), int(0.6 * INCH), int(10.8 * INCH), int(0.9 * INCH),
          "Название раздела", 28, bold=True)
    for index in range(3):
        x = int(0.6 * INCH) + index * int(3.7 * INCH)
        _card(slide, x, int(2 * INCH), int(3.4 * INCH), int(3.8 * INCH))
        _picture(slide, icon, x + int(0.2 * INCH), int(2.2 * INCH), int(0.6 * INCH),
                 int(0.6 * INCH))
        _text(slide, x + int(0.2 * INCH), int(3 * INCH), 3 * INCH, int(2.6 * INCH),
              "Опишите шаг", 14)

    # 5. Текст слева, фото справа.
    slide = prs.slides.add_slide(blank)
    _text(slide, int(0.6 * INCH), int(0.6 * INCH), int(6 * INCH), int(0.9 * INCH),
          "Название раздела", 28, bold=True)
    _list(slide, int(0.6 * INCH), int(3.1 * INCH), int(6 * INCH), int(3.1 * INCH), 16)
    _picture(slide, _photo(3, (600, 700)), int(7.2 * INCH), 0, int(4.8 * INCH), H)

    # 6. Без фото: заголовок и абзац.
    slide = prs.slides.add_slide(blank)
    _text(slide, int(0.6 * INCH), int(0.6 * INCH), int(10.8 * INCH), int(0.9 * INCH),
          "Название раздела", 28, bold=True)
    _list(slide, int(0.6 * INCH), int(3.1 * INCH), int(10.8 * INCH), int(3.1 * INCH), 16)

    # 7. Финал: декор и логотип.
    slide = prs.slides.add_slide(blank)
    _picture(slide, _flat((800, 600), "zigzag"), 7 * INCH, int(2.5 * INCH), 5 * INCH, 4 * INCH)
    _picture(slide, logo, int(0.5 * INCH), int(5.6 * INCH), int(0.8 * INCH), int(0.8 * INCH))
    _text(slide, int(0.6 * INCH), int(2.2 * INCH), 6 * INCH, int(1.6 * INCH),
          "Удачи на экзамене", 40, bold=True)

    prs.save(str(path))
    return path


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "photo_template.pptx")
