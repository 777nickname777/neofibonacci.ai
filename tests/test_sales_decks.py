"""Колода о продажах кофеен на двух шаблонах: «Дорожная карта» и шаблон с фото.

Приёмка сентября 2026 по семи замечаниям: обложка без маршрута, список в
одной карточке, пустой подзаголовок, ни одной диаграммы, белый фон вместо
фона донора, время без разбивки, фото чужой темы в колоде. План и пакет —
записанный живой прогон `sales × dorozhnaya_karta` (`llm-probe.yml`).
"""

from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

from deckwright.config import load_config
from deckwright.llm.fake import RecordedClient
from deckwright.parse.opener import parse_template
from deckwright.pipeline import STAGES, run_variant, stage_times
from deckwright.schemas import (
    ContentBlock,
    ContentPack,
    RunManifest,
    SlotRole,
    StageTiming,
)

sys.path.insert(0, str(Path(__file__).parent / "fixtures"))
from make_photo_template import build as build_photo_template

ROOT = Path(__file__).resolve().parents[1]
CONFIG = ROOT / "configs" / "config.yaml"
ROADMAP = ROOT / "data" / "holdout" / "dorozhnaya_karta.pptx"
SALES = Path(__file__).parent / "fixtures" / "roadmap_sales"
VARIANTS = ("dense", "balanced", "airy")


def _pack() -> ContentPack:
    return ContentPack.model_validate(json.loads((SALES / "pack.json").read_text("utf-8")))


def _deck(template: Path, out: Path, variant: str, recorded: Path = SALES / "recorded"):
    return run_variant(
        template_path=template,
        pack=_pack(),
        cfg=load_config(CONFIG),
        client=RecordedClient(recorded),
        variant=variant,
        output_dir=out / variant,
    )


@pytest.fixture(scope="module")
def roadmap_decks(tmp_path_factory):
    if not ROADMAP.exists():
        pytest.skip("нет бланка «Дорожная карта»")
    out = tmp_path_factory.mktemp("sales-roadmap")
    return {variant: _deck(ROADMAP, out, variant) for variant in VARIANTS}


@pytest.fixture(scope="module")
def photo_template(tmp_path_factory) -> Path:
    return build_photo_template(tmp_path_factory.mktemp("photo") / "photo.pptx")


# ── 1. Обложка и финал сохраняют оформление ─────────────────────────────────


@pytest.mark.parametrize("variant", VARIANTS)
def test_cover_keeps_route_and_project_frame(roadmap_decks, variant):
    """Маршрут из шести узлов и рамка «Название проекта» остаются на обложке."""
    template = Presentation(str(ROADMAP)).slides[0]
    nodes = [s for s in template.shapes if s.name.endswith("-node")]
    frame = next(s for s in template.shapes if s.name == "project-box")
    cover = Presentation(str(roadmap_decks[variant].pptx)).slides[0]
    boxes = {(s.left, s.top, s.width, s.height) for s in cover.shapes}
    kept = [n for n in nodes if (n.left, n.top, n.width, n.height) in boxes]
    # Разделы колоды подписывают узлы и укорачивают маршрут до своего
    # числа; не подписанный ничем маршрут остаётся целиком.
    assert len(kept) >= 4, (variant, len(kept))
    assert (frame.left, frame.top, frame.width, frame.height) in boxes
    numbers = [
        s.text_frame.text for s in cover.shapes if s.has_text_frame and s.text_frame.text.isdigit()
    ]
    assert len(numbers) == len(kept), "у каждого узла свой номер"


def test_closing_keeps_its_dark_background_and_drops_empty_cards(roadmap_decks):
    closing = Presentation(str(roadmap_decks["dense"].pptx)).slides[-1]
    assert closing._element.cSld.bg is not None
    assert roadmap_decks["dense"].deck.slides[-1].is_dark
    # Карточки итога без содержания — не оформление: пустые уходят.
    cards = [s for s in closing.shapes if s.width > 2_000_000 and s.height > 1_500_000]
    assert cards == []


# ── 2. Список — по отдельным местам, без одинаковых слайдов ─────────────────


def _headed_plan(tmp_path: Path) -> Path:
    """Живой план, где списки с заголовком — как в колоде с сайта."""
    plan = json.loads((SALES / "recorded" / "plan_deck.json").read_text("utf-8"))
    headed = {
        2: (
            "Ключевые метрики",
            ["Выручка: 47,3 млн руб.", "Чеков: 134 150", "Средний чек: 353 руб."],
        ),
        7: ("Проблемная зона", ["Выручка Мурманска 6,2 млн руб.", "Выполнение плана 87%"]),
        9: ("Приоритеты IV квартала", [
            "Запуск программы лояльности", "Удержание базы гостей", "Новые условия для B2B"
        ]),
    }
    for index, (heading, items) in headed.items():
        slide = plan["slides"][index]
        slide["subtitle"] = ""
        slide["blocks"] = [
            {"id": f"h{index}", "kind": "bullets", "heading": heading, "items": items}
        ]
    recorded = tmp_path / "recorded"
    recorded.mkdir()
    (recorded / "plan_deck.json").write_text(json.dumps(plan, ensure_ascii=False), "utf-8")
    return recorded


@pytest.mark.parametrize("variant", VARIANTS)
def test_headed_list_is_spread_and_its_heading_is_the_subtitle(tmp_path, variant):
    """«Ключевые метрики» и три пункта — не одним полем первой карточки."""
    if not ROADMAP.exists():
        pytest.skip("нет бланка «Дорожная карта»")
    result = _deck(ROADMAP, tmp_path, variant, _headed_plan(tmp_path))
    for slide in result.deck.slides:
        for element in slide.elements:
            if element.text is None or element.role is SlotRole.QUOTE:
                continue
            texts = [p.text for p in element.text.paragraphs]
            assert "Ключевые метрики" not in texts[1:] and len(texts) < 4, (
                slide.index,
                element.id,
                texts,
            )
    subtitles = [
        e.text.paragraphs[0].text
        for s in result.deck.slides
        for e in s.elements
        if e.role is SlotRole.SUBTITLE
    ]
    assert "Ключевые метрики" in subtitles


@pytest.mark.parametrize("variant", VARIANTS)
def test_no_composition_repeats_more_than_three_times(roadmap_decks, variant):
    """Пять одинаковых слайдов в колоде — против разнообразия."""
    result = roadmap_decks[variant]
    counts = Counter(
        s.pattern_id
        for s in result.deck.slides
        if s.pattern_id not in result.spec.bookend_ids
    )
    assert max(counts.values()) <= 3, counts


# ── 3. Подзаголовок — заполнен или убран ────────────────────────────────────


@pytest.mark.parametrize("variant", VARIANTS)
def test_subtitle_is_filled_or_removed(roadmap_decks, variant):
    result = roadmap_decks[variant]
    patterns = {p.id: p for p in result.spec.patterns}
    deck = Presentation(str(result.pptx))
    filled = 0
    for slide, slide_ir in zip(deck.slides, result.deck.slides, strict=True):
        pattern = patterns[slide_ir.pattern_id]
        boxes = [s.box for s in pattern.slots if s.role is SlotRole.SUBTITLE]
        for shape in slide.shapes:
            box = (shape.left, shape.top, shape.width, shape.height)
            if any(box == (b.x, b.y, b.w, b.h) for b in boxes):
                assert shape.text_frame.text.strip(), (slide_ir.index, "пустой подзаголовок")
                filled += 1
    assert filled >= 5


# ── 4. Ряды — нативными диаграммами ─────────────────────────────────────────


@pytest.mark.parametrize("variant", VARIANTS)
def test_series_become_native_charts(roadmap_decks, variant):
    """Динамика по месяцам — столбцами, доли каналов — кольцом; в каждом варианте."""
    kinds = Counter()
    for slide in Presentation(str(roadmap_decks[variant].pptx)).slides:
        for shape in slide.shapes:
            if shape.has_chart:
                kinds[shape.chart.chart_type] += 1
    assert kinds[XL_CHART_TYPE.DOUGHNUT] == 1, kinds
    assert kinds[XL_CHART_TYPE.COLUMN_CLUSTERED] >= 3, kinds


def test_table_restating_a_series_is_drawn_as_that_series(roadmap_decks):
    """Таблица «Канал / Выручка» плана — ряд «Выручка по каналам» пакета."""
    charts = [
        e.chart
        for s in roadmap_decks["dense"].deck.slides
        for e in s.elements
        if e.chart is not None and e.chart.chart_kind.value == "doughnut"
    ]
    assert charts and charts[0].series[0].values == [30.7, 8.5, 4.7, 3.4]
    assert len({c.rgb for c in charts[0].point_colors}) > 1, "у каждой доли свой цвет"


# ── 5. Фон — в test_blank_templates / test_audit ────────────────────────────
# ── 6. Время по этапам ──────────────────────────────────────────────────────


def test_stage_times_split_the_run_into_five_stages():
    def manifest(**stages: float) -> RunManifest:
        return RunManifest(
            run_id="r",
            started_at="2026-09-27T00:00:00Z",
            template_sha256="0" * 64,
            template_name="t.pptx",
            seed=1,
            variants=["v"],
            fix_mode="auto",
            timings=[StageTiming(stage=k, seconds=v) for k, v in stages.items()],
        )

    first = manifest(parse=2.0, plan=40.0, layout=0.5, render_pptx=2, render_pdf=3, audit=90)
    second = manifest(parse=0.0, plan=0.0, layout=0.5, render_pptx=2, render_pdf=5, audit=95)
    times = stage_times(44.5, [first, second])
    assert tuple(times) == STAGES
    assert times == {
        "разбор шаблона": 2.0,
        "разбор входа": 44.5,
        "план": 40.0,
        "генерация вариантов": 8.0,
        "аудит": 95.0,
    }


def test_run_summary_carries_stage_times(roadmap_decks):
    stages = stage_times(0.0, [r.manifest for r in roadmap_decks.values()])
    assert set(stages) == set(STAGES)
    # Разбор шаблона — первый вариант, остальные берут кэш.
    assert stages["разбор шаблона"] < 30


def test_plan_block_text_field_is_an_item():
    """Модель пишет абзац в `text`: живой план из-за этого уходил на повтор (43 с)."""
    block = ContentBlock.model_validate({"id": "b1", "kind": "paragraph", "text": "Рост"})
    assert block.items == ["Рост"]


def test_empty_plan_block_is_dropped_not_retried():
    """Пустой блок на разделителе живого плана `vk_tech` стоил повтора (42 с)."""
    from deckwright.schemas import SlidePlan

    slide = SlidePlan.model_validate(
        {
            "index": 9,
            "intent": "section",
            "takeaway_title": "Планы на IV квартал",
            "blocks": [{"id": "b14", "kind": "paragraph"}],
        }
    )
    assert slide.blocks == []


def test_cut_off_plan_is_an_invalid_answer():
    """План из двух слайдов при «от 10 до 15» — обрыв, а не план (run 29)."""
    from pydantic import ValidationError

    from deckwright.plan.planner import _bounded

    plan = json.loads((SALES / "recorded" / "plan_deck.json").read_text("utf-8"))
    schema = _bounded(10)
    assert len(schema.model_validate(plan).slides) == 13
    plan["slides"] = plan["slides"][:2]
    with pytest.raises(ValidationError, match="план оборван"):
        schema.model_validate(plan)


# ── 7. Фото донора — место под картинку, а не содержание ────────────────────


def test_photo_slots_are_read_from_pictures_not_names(photo_template):
    spec = parse_template(photo_template)
    with_photo = sorted(p.donor_slide_index for p in spec.patterns if p.photo_slots)
    assert with_photo == [2, 3, 5]


def test_photo_slots_are_empty_on_templates_with_brand_art():
    """3D-шар `vk_tech`, линии `vk_workspace`, маршрут бланка — не фото."""
    for name in ("data/templates/vk_tech.pptx", "data/holdout/dorozhnaya_karta.pptx"):
        path = ROOT / name
        if not path.exists():
            continue
        assert not any(p.photo_slots for p in parse_template(path).patterns), name


@pytest.mark.parametrize("variant", VARIANTS)
def test_deck_has_no_donor_photos_but_keeps_logo_and_decor(photo_template, tmp_path, variant):
    from deckwright.audit.deterministic.content import donor_photos

    result = _deck(photo_template, tmp_path, variant)
    assert donor_photos(result.pptx, result.deck) == []
    assert not [i for i in result.audit.issues if i.check_id == "integrity.donor_photo_left"]
    deck = Presentation(str(result.pptx))
    cover_pictures = [s for s in deck.slides[0].shapes if s.shape_type == 13]
    closing_pictures = [s for s in deck.slides[-1].shapes if s.shape_type == 13]
    assert len(cover_pictures) == 2 and len(closing_pictures) == 2, "логотип и декор на месте"
    icons = sum(
        1 for slide in deck.slides for s in slide.shapes if s.shape_type == 13 and s.width < 914_400
    )
    assert icons >= 2


def test_side_photo_composition_is_rebuilt_for_text(photo_template):
    """Фото сбоку уходит, заголовок и текст растягиваются на его место."""
    from deckwright.layout.matcher import _without_photos

    spec = parse_template(photo_template)
    pattern = next(p for p in spec.patterns if p.donor_slide_index == 2)
    rebuilt = _without_photos(pattern, spec)
    title = next(s for s in rebuilt.slots if s.role is SlotRole.TITLE)
    before = next(s for s in pattern.slots if s.role is SlotRole.TITLE)
    assert title.box.x < before.box.x and title.box.right == before.box.right
    assert not [s for s in rebuilt.slots if s.role is SlotRole.IMAGE]
    assert rebuilt.photo_slots == pattern.photo_slots, "место под картинку сохранено"


def test_contextual_prompt_knows_the_deck_topic():
    """Картиночный вопрос видит тему колоды: без неё чужое фото проходило «да»."""
    from deckwright.plan.planner import load_prompt

    prompt = load_prompt("audit_slide.v3")
    assert "{topic}" in prompt.template
    assert "content.visuals_on_topic" in prompt.template
